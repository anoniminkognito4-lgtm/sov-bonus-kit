"""Стенд «шаблон калькулятора как эталон».

Берёт условный тестовый кейс (tests/cases/Txx.json), заполняет им копию обезличенного шаблона
(reference/template/bonus_template_empty.xlsx), пересчитывает книгу в LibreOffice и выгружает
результаты в JSON. Два варианта:

  template — шаблон как есть (текущая практика калькуляторов);
  spec     — шаблон с правками spec_patches.py (утверждённые правила ТЗ, раздел 11).

Эталоном для приложения служит вариант spec. Вариант template нужен, чтобы видеть, где и на сколько
приложение будет отличаться от действующих калькуляторов.

Результат пишется в формате приложения (contract.py, docs/DATA_FORMATS.md раздел 2):
  tests/golden/spec/Txx.json, tests/golden/template/Txx.json — эталоны;
  reference/examples/Txx_{spec,template}.xlsx — заполненные и пересчитанные копии шаблона (их можно открыть);
  build/harness/raw/ — сырые значения ячеек, для отладки.

Запуск из корня комплекта:
    python3 tools/template_harness/fill_and_recalc.py                 # все кейсы, оба варианта
    python3 tools/template_harness/fill_and_recalc.py --case T05     # один кейс

Нужен LibreOffice (macOS: brew install --cask libreoffice) и openpyxl.
Реальные данные сюда не подставлять: только условные кейсы.
"""
import argparse
import datetime as dt
import glob
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
KIT = os.path.abspath(os.path.join(HERE, '..', '..'))
sys.path.insert(0, HERE)
import contract  # noqa: E402
import spec_patches  # noqa: E402

TEMPLATE = os.path.join(KIT, 'reference', 'template', 'bonus_template_empty.xlsx')
CALC = 'калькулятор'
INTEREST = 'проценты на отложенную часть'
FIRST_ROW, LAST_ROW, FINAL_ROW = 16, 35, 36          # промежуточные вехи 16..35, ключевая 36
PERIOD_ROW = 45                                       # таблица периодов БГБ / загрузки: 45..79
INTEREST_BLOCK = 34                                   # блок сотрудника на листе процентов: 39, 73, 107...


def soffice_path():
    for p in ('/Applications/LibreOffice.app/Contents/MacOS/soffice', shutil.which('soffice') or '',
              shutil.which('libreoffice') or '', r'C:\Program Files\LibreOffice\program\soffice.exe'):
        if p and os.path.exists(p):
            return p
    sys.exit('Не найден LibreOffice. macOS: brew install --cask libreoffice')


def d(s):
    return dt.datetime.fromisoformat(s) if s else None


def block_cols(b):
    bc = 11 + 5 * b        # K для первого сотрудника, P для второго
    return dict(ind=L(bc), weight=L(bc + 2), amount=L(bc + 3), main=L(bc + 4),
                p_from=L(bc), p_to=L(bc + 1), p_load=L(bc + 2), p_bgb=L(bc + 3), p_pb=L(bc + 4), label=L(bc))


def fill(case, variant):
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb[CALC]
    ws['B1'] = case.get('project', 'Тестовый проект')
    ws['F1'] = 'Да' if case.get('continues_from_prev_stage') else 'Нет'
    ws['F2'] = d(case.get('new_stage_start'))
    bs = case.get('bs') or {}
    ws['F3'] = 'Да' if bs.get('enabled') else 'Нет'
    ws['G3'] = bs.get('share_transferred', 1)

    ms = case['milestones']
    assert len(ms) <= LAST_ROW - FIRST_ROW + 1
    for i in range(LAST_ROW - FIRST_ROW + 1):
        r = FIRST_ROW + i
        ws['C%d' % r] = case['stage']
        if i < len(ms):
            m = ms[i]
            ws['E%d' % r] = m['name']
            ws['F%d' % r] = m.get('building') or None
            ws['G%d' % r] = d(m['plan'])
            ws['H%d' % r] = d(m.get('fact'))
        else:
            ws['E%d' % r] = None
            ws['F%d' % r] = None
            ws['G%d' % r] = None
            ws['H%d' % r] = None
    fin = case['final']
    ws['C36'] = case['stage']
    ws['E36'] = fin['name']
    ws['G36'] = d(fin['plan'])
    ws['H36'] = d(fin.get('fact'))

    names = [m['name'] for m in ms]
    emps = case['employees']
    assert len(emps) <= 2, 'в шаблоне два блока сотрудников'
    for b in range(2):
        c = block_cols(b)
        X = c['main']
        if b >= len(emps):
            for r in (1, 2, 3):
                ws['%s%d' % (X, r)] = None
            continue
        e = emps[b]
        ws['%s1' % X] = e['category']
        ws['%s2' % X] = e.get('position', '')
        ws['%s3' % X] = e['fio']
        ws['%s4' % X] = e.get('npv_start')
        ws['%s5' % X] = d(e['start'])
        ws['%s11' % X] = e.get('npv_current')
        if e.get('leave_date'):
            ws['%s12' % c['label']] = 'Дата ухода с проекта (этапа)'
            ws['%s12' % X] = d(e['leave_date'])
        if case['stage'] == 5:
            # в действующих калькуляторах 5 этапа доля авансов выставлена вручную = 33%
            ws['%s8' % X] = '=%s$7*0.33' % X
        for j, p in enumerate(e['periods']):
            r = PERIOD_ROW + j
            ws['%s%d' % (c['p_from'], r)] = d(p['from'])
            ws['%s%d' % (c['p_to'], r)] = d(p['to'])
            ws['%s%d' % (c['p_load'], r)] = p['load']
            ws['%s%d' % (c['p_bgb'], r)] = p['bgb']
            ws['%s%d' % (c['p_pb'], r)] = '=%s%d/365*(%s%d-%s%d+1)*%s%d' % (
                c['p_bgb'], r, c['p_to'], r, c['p_from'], r, c['p_load'], r)
        for k, v in (e.get('individual_dates') or {}).items():
            r = FINAL_ROW if k == 'final' else FIRST_ROW + names.index(k)
            ws['%s%d' % (c['ind'], r)] = d(v)
        for k, v in (e.get('kpi') or {}).items():
            ws['%s%d' % (X, FIRST_ROW + names.index(k))] = v
        cp = e.get('npv_checkpoints') or {}
        for r, key in ((37, 'final'), (38, 'q1'), (39, 'q2')):
            if key in cp:
                ws['%s%d' % (X, r)] = cp[key]
        # ставка ЦБ для процентов на отложенную часть (средневзвешенная за период, вводится Группой СОВ)
        wi = wb[INTEREST]
        r0 = 39 + INTEREST_BLOCK * b
        rates = case.get('key_rates') or {}
        for j in range(20):
            nm = names[j] if j < len(names) else None
            wi['I%d' % (r0 + 7 + j)] = rates.get(nm, case.get('key_rate', 0.16)) if nm else None
    stats = spec_patches.apply(wb, stage5=case['stage'] == 5) if variant == 'spec' else {}
    return wb, stats


def recalc(files, outdir):
    os.makedirs(outdir, exist_ok=True)
    prof = tempfile.mkdtemp()
    cmd = [soffice_path(), '-env:UserInstallation=file://' + prof, '--headless', '--calc',
           '--convert-to', 'xlsx', '--outdir', outdir] + files
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    shutil.rmtree(prof, ignore_errors=True)


def norm(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, float):
        if math.isnan(v):
            return None
        return v              # без округления: коэффициенты дальше умножаются на суммы
    if v == '':
        return None
    return v


def as_date(v):
    """LibreOffice иногда отдаёт даты числом (серийный номер Excel)."""
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=v)).date().isoformat()
    return norm(v)


def extract(path, case):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, wi = wb[CALC], wb[INTEREST]
    n = len(case['milestones'])
    out = {'id': case['id'], 'employees': []}
    for b, e in enumerate(case['employees']):
        c = block_cols(b)
        X, A, W = c['main'], c['amount'], c['weight']
        g = lambda ref: norm(ws[ref].value)  # noqa: E731
        emp = {'fio': e['fio'], 'category': e['category'],
               'duration_days': g('%s6' % X), 'pb': g('%s7' % X), 'advance_pool': g('%s8' % X),
               'final_pool': g('%s9' % X), 'k_fin': g('%s10' % X), 'leave_remainder': g('%s13' % X),
               'start': as_date(ws['%s5' % X].value), 'leave_date': as_date(ws['%s12' % X].value),
               'milestones': [], 'final': None, 'payouts': [], 'interest': None}
        for j in range(n):
            r = FIRST_ROW + j
            emp['milestones'].append({
                'name': ws['E%d' % r].value, 'plan': as_date(ws['G%d' % r].value), 'fact': as_date(ws['H%d' % r].value),
                'individual_plan': as_date(ws['%s%d' % (c['ind'], r)].value),
                'deviation_days': g('I%d' % r), 'pay_date': as_date(ws['J%d' % r].value),
                'weight': g('%s%d' % (W, r)), 'amount': g('%s%d' % (A, r)), 'note': g('%s%d' % (X, r))})
        emp['final'] = {'plan': as_date(ws['G36'].value), 'fact': as_date(ws['H36'].value),
                        'individual_plan': as_date(ws['%s36' % c['ind']].value),
                        'deviation_days': g('I36'), 'pay_date': as_date(ws['J36'].value), 'amount': g('%s36' % A)}
        for r in (37, 38, 39):
            emp['payouts'].append({'label': ws['E%d' % r].value, 'event_date': as_date(ws['H%d' % r].value),
                                   'pay_date': as_date(ws['J%d' % r].value), 'amount': g('%s%d' % (A, r)),
                                   'npv_checkpoint': g('%s%d' % (X, r))})
        r0 = 39 + INTEREST_BLOCK * b
        rows = []
        for j in range(n):
            r = r0 + 7 + j
            rows.append({'milestone': wi['B%d' % r].value, 'advance': norm(wi['C%d' % r].value),
                         'pay_date': as_date(wi['D%d' % r].value), 'deferred': norm(wi['E%d' % r].value),
                         'rate': norm(wi['I%d' % r].value),
                         'days_1': norm(wi['J%d' % r].value), 'interest_1': norm(wi['K%d' % r].value),
                         'days_2': norm(wi['M%d' % r].value), 'interest_2': norm(wi['N%d' % r].value),
                         'days_3': norm(wi['P%d' % r].value), 'interest_3': norm(wi['Q%d' % r].value)})
        summary = None
        for r in range(6, 38):
            if wi['B%d' % r].value == e['fio']:
                summary = {k: norm(wi['%s%d' % (col, r)].value)
                           for k, col in (('rupk_mp_total', 'D'), ('kupk_50', 'E'), ('kupk_25_1', 'F'),
                                          ('kupk_25_2', 'G'), ('total', 'H'))}
        emp['interest'] = {'rows': rows, 'summary': summary}
        out['employees'].append(emp)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--case', help='ID кейса, например T05')
    ap.add_argument('--variant', choices=['template', 'spec', 'both'], default='both')
    ap.add_argument('--cases', default=os.path.join(KIT, 'tests', 'cases'))
    ap.add_argument('--build', default=os.path.join(KIT, 'build', 'harness'))
    ap.add_argument('--golden', default=os.path.join(KIT, 'tests', 'golden'))
    ap.add_argument('--examples', default=os.path.join(KIT, 'reference', 'examples'))
    a = ap.parse_args()
    variants = ['template', 'spec'] if a.variant == 'both' else [a.variant]
    paths = sorted(glob.glob(os.path.join(a.cases, '*.json')))
    cases = [json.load(open(p, encoding='utf-8')) for p in paths]
    if a.case:
        cases = [c for c in cases if c['id'] == a.case]
    cases = [c for c in cases if c.get('harness', True)]
    filled_dir = os.path.join(a.build, 'filled')
    os.makedirs(filled_dir, exist_ok=True)
    jobs = []
    for c in cases:
        for v in variants:
            wb, stats = fill(c, v)
            p = os.path.join(filled_dir, '%s_%s.xlsx' % (c['id'], v))
            wb.save(p)
            jobs.append((c, v, p, stats))
    recalc_dir = os.path.join(a.build, 'recalc')
    recalc([p for _, _, p, _ in jobs], recalc_dir)
    os.makedirs(a.examples, exist_ok=True)
    raw_dir = os.path.join(a.build, 'raw')
    os.makedirs(raw_dir, exist_ok=True)
    for c, v, p, stats in jobs:
        rp = os.path.join(recalc_dir, os.path.basename(p))
        raw = extract(rp, c)
        with open(os.path.join(raw_dir, '%s_%s.json' % (c['id'], v)), 'w', encoding='utf-8') as f:
            json.dump(raw, f, ensure_ascii=False, indent=1)
        res = contract.build(raw, c, v)
        res['patches_applied'] = stats
        od = os.path.join(a.golden, v)
        os.makedirs(od, exist_ok=True)
        with open(os.path.join(od, c['id'] + '.json'), 'w', encoding='utf-8') as f:
            json.dump(res, f, ensure_ascii=False, indent=1)
        shutil.copy(rp, os.path.join(a.examples, '%s_%s.xlsx' % (c['id'], v)))
        print(c['id'], v, 'ok', ' '.join('%s=%d' % kv for kv in stats.items()),
              ' '.join(o['code'] for o in res['overrides_applied']))


if __name__ == '__main__':
    main()
