"""Правки формул шаблона, которые приводят его к утверждённым правилам (ТЗ, раздел 11).

Шаблон калькулятора — эталон текущей практики, но в нескольких местах он расходится с СОВ.
Чтобы получить эталон «по ТЗ», эти правки применяются к ЗАПОЛНЕННОЙ КОПИИ шаблона перед пересчётом.
Сам шаблон в reference/ не меняется.

Каждая правка = (код, описание, лист, функция замены текста формулы, строки, только 5 этап, колонки).
Правила, которые формулами шаблона не выразить, применяются после пересчёта в contract.py (overrides O1–O3).
"""
import re

from openpyxl.utils import column_index_from_string, get_column_letter

def _cap(f):
    # Потолок +50% для КУПК: в шаблоне без инд. срока 1 + 0,5 × К фин, с инд. сроком 1,5 × К фин.
    # По ТЗ всегда 1,5 × К фин (бонус не больше 150% бонуса за веху в срок).
    return re.sub(r'IF\(([A-Z]+)(\d+)="",1\+0\.5\*([A-Z]+)\$10,1\.5\*\3\$10\)', r'1.5*\3$10', f)

def _low_npv(f):
    # Правило низкого NPV включается по СТАРТОВОМУ NPV (строка 4), а не по текущему (строка 11). СОВ, сл. 24, пример 2.
    # Условие «стартовый меньше 500 млн или отрицательный» = стартовый < 500 млн (без ABS: отрицательный NPV
    # любого размера тоже идёт по правилу низкого NPV, отношение NPV при отрицательном старте не имеет смысла).
    return re.sub(r'ABS\(([A-Z]+)11\)<500000000', r'\g<1>4<500000000', f)

def _npv_ge(f):
    # Части 25% КУПК: платятся, если NPV НЕ НИЖЕ финального (>=), а не строго выше. СОВ, сл. 25.
    return re.sub(r'([A-Z]+)\$(38|39)>\1\$37', r'\1$\2>=\1$37', f)

def _leaver(f):
    # Остаток при уходе: если этап закрыт с опережением уже после ухода, К срока этапа = 1. СОВ v2, сл. 47.
    f = f.replace('(1+($I36*1%))', '(1+(MIN(0,$I36)*1%))')
    return re.sub(r'\(1\+\(\(([A-Z]+)36-\$H36\)\*1%\)\)', r'(1+(MIN(0,\g<1>36-$H36)*1%))', f)

def _paydate(f):
    # Дата выплаты: 10 число второго месяца после квартала; если выходной — предыдущий рабочий день.
    # (Праздники РФ здесь не учтены — в приложении нужен производственный календарь.)
    def repl(m):
        d = m.group(0)
        return '(%s-IF(WEEKDAY(%s,2)>5,WEEKDAY(%s,2)-5,0))' % (d, d, d)
    return re.sub(r'DATE\(YEAR\((H\d+)\),ROUNDUP\(MONTH\(\1\)/3,0\)\*3\+2,10\)', repl, f)

def _stage5_interest(f):
    # 5 этап: аванс = 33% ПБ, отложено 67% -> отложенная часть = аванс ÷ 33% × 67%.
    return f.replace('/67*33', '/33*67')

def _leaver_advances(f, cell):
    # Аванс за веху, закрытую после даты ухода, не платится (СОВ v2, сл. 47). В шаблоне платится.
    # Колонка суммы N (первый сотрудник) или S (второй); дата ухода — в строке 12 следующей колонки (O или T).
    # Сравнивается фактическая дата вехи (H): другой даты в шаблоне нет; в приложении — дата закрытия в 1С.
    r = cell.row
    X = get_column_letter(column_index_from_string(cell.column_letter) + 1)
    head = '=IFERROR(IF($I%d<-100,0,' % r
    if f.startswith(head) and f.endswith(',"")'):
        body = f[len(head):-len(',"")')]
        return '=IFERROR(IF(AND(%s$12<>"",$H%d>%s$12),0,IF($I%d<-100,0,%s),"")' % (X, r, X, r, body)
    return f


CALC = 'калькулятор'
INTEREST = 'проценты на отложенную часть'

# (код, описание, лист, функция, строки (None = все), только для 5 этапа, колонки (None = все))
PATCHES = [
    ('P1_cap', 'Потолок +50% у КУПК = 1,5 × К фин в обоих случаях', CALC, _cap, range(16, 36), False, None),
    ('P2_low_npv', 'Правило низкого NPV по стартовому NPV (< 500 млн, без ABS)', CALC, _low_npv, [10], False, None),
    ('P3_npv_ge', 'Части 25%: NPV не ниже финального (>=)', CALC, _npv_ge, [38, 39], False, None),
    ('P4_leaver', 'Остаток при уходе: К срока этапа не больше 1', CALC, _leaver, [13], False, None),
    ('P5_paydate', 'Дата выплаты: перенос с выходного на пятницу', CALC, _paydate, range(16, 40), False, None),
    ('P6_stage5_interest', '5 этап: отложенная часть = аванс ÷ 33% × 67%', INTEREST, _stage5_interest, None, True, None),
    ('P7_leaver_advances', 'Аванс за веху, закрытую после ухода, не платится', CALC, _leaver_advances, range(16, 36), False, ('N', 'S')),
]

def apply(wb, stage5=False):
    """Применяет правки к открытой книге openpyxl. Возвращает {код: число изменённых формул}."""
    stats = {}
    for code, _, sheet, fn, rows, only5, cols in PATCHES:
        if only5 and not stage5:
            continue
        nargs = fn.__code__.co_argcount
        ws = wb[sheet]
        n = 0
        for row in ws.iter_rows():
            for c in row:
                if rows is not None and c.row not in rows:
                    continue
                if cols is not None and c.column_letter not in cols:
                    continue
                call = (lambda t, c=c: fn(t, c)) if nargs == 2 else fn
                v = c.value
                if hasattr(v, 'text'):            # формула массива (ArrayFormula)
                    t = v.text if v.text.startswith('=') else '=' + v.text
                    nt = call(t)
                    if nt != t:
                        v.text = nt if v.text.startswith('=') else nt[1:]
                        n += 1
                elif isinstance(v, str) and v.startswith('='):
                    nv = call(v)
                    if nv != v:
                        c.value = nv
                        n += 1
        stats[code] = n
    return stats
