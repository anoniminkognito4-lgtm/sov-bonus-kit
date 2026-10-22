"""Реестр одним взглядом и калькулятор по выбранным этапам."""

from io import BytesIO

import pytest
from conftest import ROOT
from openpyxl import load_workbook

from sov.exports.calculator import calculator_workbook
from sov.registry.session import Registry
from sov.services.panel import Actor, PanelService
from sov.services.registry_view import RegistryView

AS_OF = "2026-09-30"


@pytest.fixture(scope="module")
def loaded(tmp_path_factory):
    registry = Registry(tmp_path_factory.mktemp("view") / "sov.sqlite")
    service = PanelService(registry, Actor("Тестовый методолог"))
    service.demo(ROOT / "samples")
    yield service, RegistryView(registry.rows)
    registry.dispose()


def row_for(view, person_id, as_of=AS_OF):
    return next(r for r in view.overview(as_of) if r["person_id"] == person_id)


def test_one_row_per_assignment(loaded):
    service, view = loaded
    rows = view.overview(AS_OF)
    assert len(rows) == len(service.registry.rows("assignment"))
    assert all(r["fio"] and r["project"] for r in rows)


def test_overdue_open_milestones_counted_on_cut_date(loaded):
    _, view = loaded
    row = row_for(view, "T0023")
    assert row["overdue_count"] == 4
    assert row["max_overdue_days"] == 365
    assert row["closed"] == 2
    # На более раннюю дату просрочка меньше: срез не зависит от даты компьютера.
    assert row_for(view, "T0023", "2026-01-01")["max_overdue_days"] < 365


def test_individual_dates_and_newcomer_milestones(loaded):
    _, view = loaded
    newcomer = row_for(view, "T0011")
    assert newcomer["individual_count"] == 1
    assert newcomer["total"] < 7  # вехи первых 60 дней не участвуют
    assignment = next(a for a in view.rows("assignment") if a["person_id"] == "T0011")
    states = {m["state"] for m in view.milestones_for(assignment, AS_OF)}
    assert "не участвует" in states


def test_status_and_facts_after_cut_date(loaded):
    _, view = loaded
    assert row_for(view, "T0028")["status"] == "ушёл"
    early = row_for(view, "T0001", "2025-01-01")
    assert early["status"] == "ещё не начал"
    assert early["closed"] == 0  # факт после даты среза ещё не виден


def test_calculator_for_several_stages(loaded):
    service, view = loaded
    options = view.stage_options()[:2]
    identifiers = view.assignments_for(options)
    calculation = service.calculate("2026-Q3", identifiers)
    labels = {o: view.stage_label(o) for o in options}
    data = calculator_workbook(calculation.results, options, labels, "2026-Q3")
    assert data == calculator_workbook(calculation.results, options, labels, "2026-Q3")
    book = load_workbook(BytesIO(data))
    assert len(book.sheetnames) == 2
    sheet = book[book.sheetnames[0]]
    text = [c for row in sheet.iter_rows(values_only=True) for c in row if isinstance(c, str)]
    assert text[0].startswith("Калькулятор · ")
    assert "Проектный бонус (ПБ) этапа, ₽" in text
    assert any(t.startswith("Ключевая веха · ") for t in text)
    assert all(not isinstance(c, str) or not c.startswith("=") for c in text)
