"""M6: источники → реестр → особые случаи → неизменный закрытый снимок."""

import csv
from copy import deepcopy

import pytest
from conftest import ROOT
from openpyxl import load_workbook
from test_panel import fix_sources, panel  # noqa: F401

from sov.adapters.registry_case import RegistryAdapter
from sov.exports.files import payroll
from sov.services.panel import Actor, PanelService


def csv_source(tmp_path, name, rows, fields=None):
    path = tmp_path / name
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields or list(rows[0]), delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
    return path


def sample_row(name, tab="T0001"):
    with (ROOT / "samples" / name).open(encoding="utf-8-sig", newline="") as stream:
        return next(
            dict(r) for r in csv.DictReader(stream, delimiter=";") if r.get("tab_no") == tab
        )


def apply(panel, path):  # noqa: F811
    preview = panel.preview(path)
    assert not preview.errors, preview.errors
    return panel.apply(preview)


def test_probation_release_is_recognized_in_release_quarter(panel, tmp_path):  # noqa: F811
    row = sample_row("hr_employees.csv")
    row.update(probation_end="2025-09-30", probation_result="пройден")
    apply(panel, csv_source(tmp_path, "Испытательный срок.csv", [row]))
    identifier = next(
        a["id"] for a in panel.registry.rows("assignment") if a["person_id"] == "T0001"
    )
    panel.save_decision(
        "probation_release",
        identifier,
        {"pay_date": "2025-11-10"},
        "Условное окончание ИС",
        "условное_письмо.pdf",
    )
    old = panel.calculate("2025-Q2", [identifier])
    assert not old.entries
    released = panel.calculate("2025-Q4", [identifier])
    assert len(released.entries) >= 1
    delayed = [
        r
        for r in released.entries
        if r["milestone"] in ("Мотивационная веха 1", "Мотивационная веха 2")
    ]
    assert len(delayed) == 2
    assert all(r["pay_date"] == "2025-11-10" for r in delayed)
    assert all(r["calc_status"] == "рассчитано" for r in released.entries)


def test_sale_source_needs_rp_work_end_and_uses_previous_npv(panel, tmp_path):  # noqa: F811
    with (ROOT / "samples/projects.csv").open(encoding="utf-8-sig", newline="") as stream:
        project = next(
            dict(r) for r in csv.DictReader(stream, delimiter=";") if r["project_code"] == "TP-T02"
        )
    project["sale_date"] = "2026-01-15"
    apply(panel, csv_source(tmp_path, "Продажа проекта.csv", [project]))
    identifier = next(
        a["id"] for a in panel.registry.rows("assignment") if a["person_id"] == "T0002"
    )
    assert any(
        p["code"] == "MANUAL_CALC" for p in panel.calculate("2026-Q1", [identifier]).problems
    )
    panel.save_decision(
        "sale_end",
        identifier,
        {"work_end": "2025-12-31"},
        "Условное завершение работы",
        "условное_письмо.pdf",
    )
    value = panel.calculate("2026-Q1", [identifier]).results[identifier]
    assert value["case"]["sale"]["npv_previous"] == 5_000_000_000
    assert len(value["result"]["employees"][0]["payouts"]) == 1
    assert value["result"]["employees"][0]["final"]["kind"] == "остаток при продаже"


def milestone_copy(tmp_path, end, fact=None):
    workbook = load_workbook(ROOT / "samples/milestones_1c.xlsx")
    sheet = workbook.active
    columns = {cell.value: cell.column for cell in sheet[1]}
    for row in sheet.iter_rows(min_row=2):
        if (
            row[columns["Код проекта"] - 1].value == "TP-T01"
            and row[columns["Ключевая"] - 1].value == "Да"
        ):
            row[columns["Плановый срок"] - 1].value = end
            row[columns["Фактический срок"] - 1].value = fact
            row[columns["Дата закрытия в 1С"] - 1].value = fact
            row[columns["Статус"] - 1].value = "Выполнена" if fact else "Назначена"
    path = tmp_path / f"Вехи {end}.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def test_extension_preserves_accrued_advance_and_closed_export(panel, tmp_path):  # noqa: F811
    fix_sources(panel, tmp_path)
    key = next(
        m["milestone_code"]
        for m in panel.registry.rows("milestone")
        if m["project_code"] == "TP-T01" and m["is_key"]
    )
    # MILESTONES §6.3: снять факт ключевой вехи можно только решением.
    panel.save_milestone_decision(
        "fact_change_accept", {"milestone_code": key, "fact": None}, "Условное письмо", "f.pdf"
    )
    apply(panel, milestone_copy(tmp_path, "30.09.2026"))
    identifier = next(
        a["id"] for a in panel.registry.rows("assignment") if a["person_id"] == "T0001"
    )
    old_case = panel.calculate("2026-Q3", [identifier]).results[identifier]["case"]
    old_forecast = panel.simulate(
        identifier, {"dates": {"final": "2026-09-30"}}, as_of="2026-09-30"
    )
    panel.close(panel.calculate("2026-Q3"))
    closed = panel.closed("2026-Q3")
    export = payroll(closed)
    locked = closed["payload"]["results"][identifier]["result"]["employees"][0]["milestones"][-1]
    apply(panel, milestone_copy(tmp_path, "31.12.2026"))
    value = panel.calculate("2026-Q4", [identifier]).results[identifier]
    result = value["result"]["employees"][0]
    assert value["case"]["final"]["plan"] == "2026-12-31"
    assert result["pb"] > old_forecast["baseline"]["employees"][0]["pb"]
    assert result["milestones"][-1]["amount"] == locked["amount"]
    assert result["milestones"][-1]["frozen"]
    assert value["case"]["employees"][0]["periods"][-1]["to"] == "2026-12-31"
    forecast = panel.simulate(identifier, {"dates": {"final": "2026-12-31"}}, as_of="2026-12-31")
    assert (
        forecast["scenario"]["employees"][0]["final"]["amount"]
        > old_forecast["scenario"]["employees"][0]["final"]["amount"]
    )
    assert panel.closed("2026-Q3") == closed and payroll(panel.closed("2026-Q3")) == export
    assert old_case["final"]["plan"] == "2026-09-30"


def test_transfer_creates_separate_assignment_and_requires_month_end_npv(panel, tmp_path):  # noqa: F811
    old = sample_row("assignments.csv")
    old.update(date_to="2026-03-31", end_reason="перевод")
    new = {
        **old,
        "date_from": "2026-04-01",
        "date_to": "2026-09-30",
        "end_reason": "",
        "category": "КУПК",
        "sov_type": "Вехи от NPV",
    }
    apply(panel, csv_source(tmp_path, "Перевод.csv", [old, new]))
    assignments = [a for a in panel.registry.rows("assignment") if a["person_id"] == "T0001"]
    assert len(assignments) == 2 and len({a["id"] for a in assignments}) == 2
    previous = next(a for a in assignments if a["category"] == "РУПК")
    current = next(a for a in assignments if a["category"] == "КУПК")
    assert previous["date_to"] == "2026-03-31" and current["date_from"] == "2026-04-01"
    npv = {
        "tab_no": "T0001",
        "project_code": "TP-T01",
        "stage": 4,
        "npv_start": 5_000_000_000,
        "fixed_date": "2026-04-01",
        "basis": "Условная ФМ",
    }
    assert panel.preview(csv_source(tmp_path, "Ошибочная ФМ.csv", [npv])).errors
    npv["fixed_date"] = "2026-04-30"
    apply(panel, csv_source(tmp_path, "Месячная ФМ.csv", [npv]))
    case = RegistryAdapter(panel.registry).build_case(current["id"])
    assert case["employees"][0]["npv_start"] == 5_000_000_000
    assert case["employees"][0]["start"] == "2026-04-01"
    assert (
        RegistryAdapter(panel.registry).build_case(previous["id"])["employees"][0]["category"]
        == "РУПК"
    )


def add_rupk_peer(panel, tmp_path):  # noqa: F811
    hr = sample_row("hr_employees.csv")
    hr.update(tab_no="T9001", fio="Тестов Условный Участник")
    apply(panel, csv_source(tmp_path, "Новый условный сотрудник.csv", [hr]))
    bgb = sample_row("hr_bgb_history.csv")
    bgb["tab_no"] = "T9001"
    apply(panel, csv_source(tmp_path, "Условный БГБ.csv", [bgb]))
    assignment = sample_row("assignments.csv")
    assignment["tab_no"] = "T9001"
    apply(panel, csv_source(tmp_path, "Условное назначение.csv", [assignment]))
    return [a for a in panel.registry.rows("assignment") if a["project_code"] == "TP-T01"]


def test_atomic_group_redistribution_and_revalidation_after_bgb_change(panel, tmp_path):  # noqa: F811
    team = add_rupk_peer(panel, tmp_path)
    before = panel.calculate("2025-Q2")
    name = "Мотивационная веха 1"
    rows = [
        {"assignment_id": a["id"], "milestone": name, "pct": pct}
        for a, pct in zip(team, (0.3, -0.3), strict=True)
    ]
    original_revision = panel.registry.revision()
    invalid = deepcopy(rows)
    invalid[1]["pct"] = -0.2
    with pytest.raises(ValueError):
        panel.save_redistribution(invalid, "T0003", "2025-07-01", "Условное письмо", "письмо.pdf")
    assert panel.registry.revision() == original_revision
    panel.save_redistribution(rows, "T0003", "2025-07-01", "Условное письмо", "письмо.pdf")
    after = panel.calculate("2025-Q2")
    for a, pct in zip(team, (0.3, -0.3), strict=True):
        old = before.results[a["id"]]["result"]["employees"][0]
        new = after.results[a["id"]]["result"]["employees"][0]
        assert new["milestones"][0]["amount"] == pytest.approx(
            old["milestones"][0]["amount"] * (1 + pct)
        )
        assert new["final"]["amount"] == pytest.approx(old["final"]["amount"] * (1 + pct / 6))
    bgb = sample_row("hr_bgb_history.csv")
    bgb.update(tab_no="T9001", bgb="3000000")
    apply(panel, csv_source(tmp_path, "Изменённый условный БГБ.csv", [bgb]))
    assert any(p["code"] == "MANUAL_CALC" for p in panel.calculate("2025-Q2").problems)
    with pytest.raises(PermissionError):
        PanelService(panel.registry, Actor("Сотрудник", "Сотрудник", "T0001")).save_redistribution(
            rows, "T0003", "2025-07-01", "Письмо", "письмо.pdf"
        )


def test_registry_editions_apply_to_engine_scenario_and_snapshot(panel, tmp_path):  # noqa: F811
    fix_sources(panel, tmp_path)
    rows = [
        {"edition": "Условная редакция", "name": "effective_from", "value": "2026-Q3"},
        {
            "edition": "Условная редакция",
            "name": "advance_shares",
            "value": "[0.6,0.6,0.6,0.6,0.3]",
        },
    ]
    apply(panel, csv_source(tmp_path, "Параметры СОВ.csv", rows))
    calculation = panel.calculate("2026-Q3")
    assert calculation.params.edition == "Условная редакция"
    assert all(r["edition"] == "Условная редакция" for r in calculation.entries)
    identifier = next(
        a["id"] for a in panel.registry.rows("assignment") if a["person_id"] == "T0001"
    )
    simulation = panel.simulate(identifier, {}, as_of="2026-09-30")
    assert simulation["baseline"]["edition"] == "Условная редакция"
    panel.close(calculation)
    closed = panel.closed("2026-Q3")
    assert closed["payload"]["params"]["advance_shares"][0] == 0.6
    from sov.engine.calculate import calculate_case
    from sov.engine.params import restore_params

    params = restore_params(closed["payload"]["params"])
    for value in closed["payload"]["results"].values():
        assert calculate_case(value["case"], params) == value["result"]
    old_csv = payroll(closed, "csv")
    rows[1]["value"] = "[0.5,0.5,0.5,0.5,0.3]"
    apply(panel, csv_source(tmp_path, "Новые параметры.csv", rows))
    assert payroll(panel.closed("2026-Q3"), "csv") == old_csv
    assert panel.calculate("2026-Q2").params.edition != "Условная редакция"


def test_closed_parts_are_not_paid_twice_when_probation_sources_change(panel, tmp_path):  # noqa: F811
    fix_sources(panel, tmp_path)
    panel.close(panel.calculate("2026-Q3"))
    identifier = next(
        a["id"] for a in panel.registry.rows("assignment") if a["person_id"] == "T0001"
    )
    hr = sample_row("hr_employees.csv")
    hr.update(probation_end="2026-10-31", probation_result="пройден")
    apply(panel, csv_source(tmp_path, "Новое окончание ИС.csv", [hr]))
    panel.save_decision(
        "probation_release", identifier, {"pay_date": "2026-11-10"}, "Условное письмо", "письмо.pdf"
    )
    calculation = panel.calculate("2026-Q4", [identifier])
    assert not any(
        r["part"] or r["milestone"] == "Мотивационная веха 6" for r in calculation.entries
    )
    assert calculation.results[identifier]["result"]["employees"][0]["payouts"][0]["frozen"]


def test_group_decision_remains_valid_after_accrual_and_cannot_rewrite_it(panel, tmp_path):  # noqa: F811
    team = add_rupk_peer(panel, tmp_path)
    fix_sources(panel, tmp_path)
    rows = [
        {"assignment_id": a["id"], "milestone": "Мотивационная веха 6", "pct": pct}
        for a, pct in zip(team, (0.3, -0.3), strict=True)
    ]
    panel.save_redistribution(rows, "T0003", "2026-09-30", "Условное письмо", "письмо.pdf")
    calculation = panel.calculate("2026-Q3")
    assert not calculation.blocking
    panel.close(calculation)
    current = panel.calculate("2026-Q4")
    assert not any(p["code"] == "MANUAL_CALC" for p in current.problems)
    for row in rows:
        result = current.results[row["assignment_id"]]["result"]["employees"][0]
        assert result["milestones"][-1]["frozen"]
        assert result["milestones"][-1]["redistribution_pct"] == row["pct"]
    changes = deepcopy(rows)
    changes[0]["pct"], changes[1]["pct"] = 0.2, -0.2
    with pytest.raises(ValueError, match="заморожен"):
        panel.save_redistribution(changes, "T0003", "2026-12-31", "Новое письмо", "письмо.pdf")


def test_correction_only_in_next_open_quarter_does_not_change_original(panel, tmp_path):  # noqa: F811
    fix_sources(panel, tmp_path)
    panel.close(panel.calculate("2026-Q3"))
    original = panel.closed("2026-Q3")
    export = payroll(original)
    identifier = next(
        a["id"] for a in panel.registry.rows("assignment") if a["person_id"] == "T0001"
    )
    payload = {
        "snapshot_id": original["id"],
        "quarter": "2026-Q4",
        "amount": -1000.0,
        "pay_date": "2027-02-10",
    }
    panel.save_decision("correction", identifier, payload, "Условная корректировка", "письмо.pdf")
    rows = [
        r
        for r in panel.calculate("2026-Q4").entries
        if r["kind"] == "корректировка закрытого квартала"
    ]
    assert len(rows) == 1 and rows[0]["amount"] == -1000
    assert rows[0]["correction_snapshot_id"] == original["id"]
    with pytest.raises(ValueError):
        panel.save_decision(
            "correction", identifier, {**payload, "quarter": "2026-Q3"}, "Письмо", "письмо.pdf"
        )
    assert panel.closed("2026-Q3") == original and payroll(panel.closed("2026-Q3")) == export


def test_csv_redistribution_preview_validates_full_fund(panel, tmp_path):  # noqa: F811
    add_rupk_peer(panel, tmp_path)
    rows = [
        {
            "tab_no": tab,
            "project_code": "TP-T01",
            "stage": 4,
            "milestone": "Мотивационная веха 1",
            "pct": pct,
            "decided_by": "T0003",
            "approved_date": "2025-07-01",
            "basis": "Условное письмо",
        }
        for tab, pct in (("T0001", 0.3), ("T9001", -0.2))
    ]
    invalid = panel.preview(csv_source(tmp_path, "Несбалансированное письмо.csv", rows))
    assert invalid.errors
    rows[1]["pct"] = -0.3
    apply(panel, csv_source(tmp_path, "Согласованное письмо.csv", rows))
    result = panel.calculate("2025-Q2")
    assert not any(p["code"] == "MANUAL_CALC" for p in result.problems)
