"""M3: права, прогнозы, закрытие квартала и неизменяемые выгрузки."""

from io import BytesIO

import pytest
from conftest import ROOT
from openpyxl import load_workbook
from sqlalchemy import func, select

from sov.exports.files import money, payroll, payroll_rows, statement
from sov.registry.models import audit_log
from sov.registry.session import Registry
from sov.services.panel import Actor, PanelService


def corrected_sources(folder):
    from tools.make_panel_demo import make_panel_demo

    return make_panel_demo(folder)


@pytest.fixture
def panel(tmp_path):
    registry = Registry(tmp_path / "sov.sqlite")
    service = PanelService(registry, Actor("Тестовый методолог"))
    service.demo(ROOT / "samples")
    yield service
    registry.dispose()


def fix_sources(panel, tmp_path):
    for path in corrected_sources(tmp_path):
        panel.apply(panel.preview(path))
    assignment = next(a for a in panel.registry.rows("assignment") if a["person_id"] == "T0039")
    panel.save_decision(
        "early_stage",
        assignment["id"],
        {"approved": True},
        "Условное согласование опережения",
        "условное_письмо.pdf",
    )


def test_forecasts_and_missing_data_prevent_closing(panel):
    calculation = panel.calculate("2026-Q3")
    assert calculation.blocking
    assert any(p["code"] == "FM_NOT_APPROVED" for p in calculation.blocking)
    assert any(e["calc_status"] == "прогноз" for e in calculation.entries)
    with pytest.raises(ValueError, match="блокирующие"):
        panel.close(calculation)
    assert not panel.closed("2026-Q3")
    assert not panel.registry.rows("accrual")


def test_complete_quarter_and_frozen_export(panel, tmp_path):
    fix_sources(panel, tmp_path)
    calculation = panel.calculate("2026-Q3")
    assert not calculation.blocking
    identifier = panel.close(calculation)
    closed = panel.closed("2026-Q3")
    assert closed["id"] == identifier
    original_csv = payroll(closed, "csv")
    assert "Табельный номер" in original_csv.decode("utf-8-sig")
    book = load_workbook(BytesIO(payroll(closed)), data_only=False)
    assert book["К начислению"].max_row > 1
    assert all(cell.data_type != "f" for row in book["К начислению"] for cell in row)
    book.close()
    # Последующие правки исходных данных не меняют уже закрытый квартал и его ФИО/суммы.
    assignment = panel.registry.rows("assignment")[0]
    panel.save_decision(
        "start_npv",
        assignment["id"],
        {"value": 8_000_000_000},
        "Корректировка следующего расчёта",
        "новое_решение.pdf",
    )
    assert payroll(panel.closed("2026-Q3"), "csv") == original_csv
    assert len(payroll_rows(closed)) > 0
    with pytest.raises(ValueError, match="закрыт"):
        panel.close(panel.calculate("2026-Q3"))


def test_stale_and_partial_calculations_cannot_close(panel, tmp_path):
    fix_sources(panel, tmp_path)
    first = panel.calculate("2026-Q3")
    identifier = first.assignment_ids[0]
    panel.save_decision(
        "start_npv", identifier, {"value": 5_000_000_000}, "Основание", "решение.pdf"
    )
    with pytest.raises(ValueError, match="изменились"):
        panel.close(first)
    with pytest.raises(ValueError, match="все назначения"):
        panel.close(panel.calculate("2026-Q3", [identifier]))


def test_permissions_enforced_in_services(panel):
    for role in ("Сотрудник", "Руководитель"):
        restricted = PanelService(panel.registry, Actor("Тестовый пользователь", role))
        for operation in (
            lambda: restricted.demo(ROOT / "samples"),
            lambda: restricted.preview(ROOT / "samples/key_rates.csv"),
            lambda: restricted.calculate("2026-Q3"),
            lambda: restricted.decisions(),
            lambda: restricted.closed("2026-Q3"),
            lambda: restricted.journal(),
            lambda: restricted.export_event("payroll", "test"),
            lambda: restricted.save_decision("early_stage", "test", {}, "a", "b"),
        ):
            with pytest.raises(PermissionError):
                operation()


def test_manual_decision_requires_basis_and_explicit_approval(panel):
    assignment = panel.registry.rows("assignment")[0]
    with pytest.raises(ValueError, match="основание"):
        panel.save_decision("early_stage", assignment["id"], {"approved": True}, "", "file")
    with pytest.raises(ValueError, match="подтверждение"):
        panel.save_decision("early_stage", assignment["id"], {"approved": False}, "basis", "file")
    assert not panel.decisions()


def test_outside_session_recomputes_interest(panel):
    assignment = next(a for a in panel.registry.rows("assignment") if a["person_id"] == "T0001")
    before = panel.calculate("2026-Q3", [assignment["id"]]).results[assignment["id"]]["result"][
        "employees"
    ][0]
    panel.save_decision(
        "outside_session",
        assignment["id"],
        {"milestone": "Мотивационная веха 1", "pay_date": "2025-07-25"},
        "Вне сессии",
        "письмо.pdf",
    )
    after = panel.calculate("2026-Q3", [assignment["id"]]).results[assignment["id"]]["result"][
        "employees"
    ][0]
    assert after["milestones"][0]["pay_date"] == "2025-07-25"
    assert after["interest"]["total"] > before["interest"]["total"]


def test_export_rounding_and_literal_excel_text(panel):
    assert str(money(1.005)) == "1.01"
    calculation = panel.calculate("2026-Q3")
    results = calculation.results
    value = next(iter(results.values()))
    value["result"]["employees"][0]["fio"] = '=HYPERLINK("x")'
    book = load_workbook(
        BytesIO(statement(results, value["case"]["project_code"], value["case"]["stage"]))
    )
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    book.close()


def test_view_and_export_actions_are_audited(panel):
    with panel.registry.engine.connect() as connection:
        before = connection.scalar(select(func.count()).select_from(audit_log))
    panel.view("employee", "T0001")
    panel.export_event("statement", "TP-T01")
    with panel.registry.engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(audit_log)) == before + 2


def test_amount_change_control_compares_same_quarter(panel):
    before = panel.calculate("2026-Q3")
    assignment = next(a for a in panel.registry.rows("assignment") if a["person_id"] == "T0002")
    panel.save_decision(
        "start_npv",
        assignment["id"],
        {"value": 10_000_000_000},
        "Условная корректировка",
        "решение.pdf",
    )
    after = panel.calculate("2026-Q3", previous=before)
    assert any(p["code"] == "AMOUNT_CHANGE_10PCT" and not p["blocking"] for p in after.problems)
    other_quarter = panel.calculate("2026-Q2", previous=before)
    assert not any(p["code"] == "AMOUNT_CHANGE_10PCT" for p in other_quarter.problems)


def test_close_is_atomic_on_storage_error(panel, tmp_path, monkeypatch):
    import sov.services.panel as module

    fix_sources(panel, tmp_path)
    calculation = panel.calculate("2026-Q3")
    before = panel.registry.domain_state()
    original = module._domain_insert

    def fail_after_insert(*args, **kwargs):
        original(*args, **kwargs)
        raise RuntimeError("Условная ошибка закрытия")

    monkeypatch.setattr(module, "_domain_insert", fail_after_insert)
    with pytest.raises(RuntimeError):
        panel.close(calculation)
    assert panel.registry.domain_state() == before
    assert not panel.closed("2026-Q3")
    assert panel.registry.revision() == calculation.revision


def test_expense_decision_applies_to_project_stage(panel):
    assignment = next(a for a in panel.registry.rows("assignment") if a["person_id"] == "T0037")
    panel.save_decision(
        "stage5_expense",
        assignment["id"],
        {"plan_budget": 100_000_000, "fact_budget": 150_000_000},
        "Новая ФМ",
        "ФМ.pdf",
    )
    calculation = panel.calculate("2026-Q3")
    affected = [
        v["case"] for v in calculation.results.values() if v["case"]["project_code"] == "TP-T17b"
    ]
    assert len(affected) == 2
    assert all(case["stage5_expense"]["fact_budget"] == 150_000_000 for case in affected)
