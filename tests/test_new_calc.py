"""M9: новый калькулятор, сценарии C01–C14 (docs/NEW_CALC.md §9)."""

import re
import zipfile

import pytest
from conftest import ROOT
from openpyxl import load_workbook

from sov.exports.calc_file import verify, write_calculator_file
from sov.loaders.service import apply_preview, apply_records, prepare_upload, rollback_upload
from sov.registry.session import Registry
from sov.services import new_calc
from sov.services.panel import Actor, PanelService

TEMPLATE = ROOT / "reference/template/bonus_template_empty.xlsx"
FOLDER = ROOT / "samples/milestones"


@pytest.fixture
def panel(tmp_path):
    registry = Registry(tmp_path / "sov.sqlite")
    apply_preview(registry, prepare_upload(registry, FOLDER / "ms_projects.csv"))
    apply_preview(
        registry,
        prepare_upload(
            registry, FOLDER / "ms_seed.xlsx", options={"export_date": "2026-07-01", "mode": "seed"}
        ),
    )
    yield PanelService(registry, Actor("Тестовый методолог"))
    registry.dispose()


def add_milestone(panel, code, project, queue, stage, name, plan, key=False, **extra):
    records = [
        (
            "milestone",
            {
                "milestone_code": code,
                "project_code": project,
                "queue": queue,
                "stage": stage,
                "name": name,
                "is_key": int(key),
                "status": extra.pop("status", "Назначена"),
                "queue_raw": extra.pop("queue_raw", queue),
                **extra,
            },
        )
    ]
    if plan:
        records.append(
            (
                "milestone_plan_version",
                {"milestone_code": code, "version": 1, "plan_date": plan, "basis": "тест"},
            )
        )
    apply_records(panel.registry, records, f"тест {code}", "тест", "test")


def selection(panel, project="TP-M1", queue="1", stage=4):
    return new_calc.select_milestones(new_calc.Store(panel.registry.rows), project, queue, stage)


def names(items):
    return [m["name"] for m in items]


def codes(result):
    return {p["code"] for p in result.problems}


def test_c01_queue_and_project_milestones_not_after_key(panel):
    add_milestone(panel, "X-LATE", "TP-M1", "", 4, "Веха проекта поздняя", "2027-03-01")
    result = selection(panel)
    assert result.key["name"] == "Получено РнВ"
    assert "Общая веха проекта" in names(result.milestones)  # без очереди, раньше ключевой
    assert "Веха проекта поздняя" not in names(result.milestones)
    assert ("Веха проекта поздняя", "веха проекта позже ключевой") in [
        (m["name"], reason) for m, reason in result.excluded
    ]
    assert "Проект завершен" not in names(result.milestones)  # другой этап
    assert "Веха А" in names(result.milestones) and len(result.milestones) == 5


def test_c02_two_keys_earliest_and_question(panel):
    result = selection(panel, stage=5)
    assert result.key["name"] == "Проект завершен"
    assert "NEWCALC_MULTI_KEY" in codes(result) and not result.blocking


def test_c03_no_key_blocks_nothing_created(panel):
    result = selection(panel, queue="2")
    assert "NEWCALC_NO_KEY" in codes(result) and result.blocking
    with pytest.raises(ValueError, match="ключевой"):
        panel.create_calculator("TP-M1", "2", 4, {"start_date": "2026-01-01"}, [])
    assert not panel.registry.rows("stage") and not panel.registry.rows("calc_milestone")


def test_c04_sorted_by_date_after_key_flagged(panel):
    add_milestone(panel, "X-AFTER", "TP-M1", "1", 4, "Веха после ключевой", "2027-02-01")
    result = selection(panel)
    plans = [
        new_calc.Store(panel.registry.rows).plan(m["milestone_code"]) for m in result.milestones
    ]
    assert plans == sorted(plans)
    assert result.milestones[-1]["name"] == "Веха после ключевой"
    assert "NEWCALC_AFTER_KEY" in codes(result)


def test_c05_no_plan_and_inactive_excluded_with_reason(panel):
    add_milestone(panel, "X-NOPLAN", "TP-M1", "1", 4, "Веха без даты", None)
    add_milestone(panel, "X-CANCEL", "TP-M1", "1", 4, "Отменённая", "2026-08-01", status="Отменена")
    code = next(
        m["milestone_code"] for m in panel.registry.rows("milestone") if m["name"] == "Веха Б"
    )
    panel.save_milestone_decision("milestone_inactive", {"milestone_code": code}, "письмо", "f.pdf")
    result = selection(panel)
    reasons = {m["name"]: reason for m, reason in result.excluded}
    assert reasons["Веха без даты"] == "нет плановой даты"
    assert reasons["Веха Б"] == "снята"
    assert reasons["Отменённая"].startswith("неизвестный статус 1С")
    assert not {"Веха без даты", "Веха Б", "Отменённая"} & set(names(result.milestones))


def test_c06_queue_text_differs_by_spaces(panel):
    add_milestone(
        panel, "X-SPACE", "TP-M1", "1", 4, "Веха с пробелом", "2026-08-15", queue_raw="1 "
    )
    assert "NEWCALC_QUEUE_TEXT" in codes(selection(panel))


def create(panel, project="TP-M1", queue="1", stage=4, start="2025-12-01", team=()):
    return panel.create_calculator(project, queue, stage, {"start_date": start}, list(team))


def test_c07_dates_follow_store_new_milestone_needs_decision(panel):
    create(panel)
    store = new_calc.Store(panel.registry.rows)
    assert not new_calc.set_changes(store)
    preview = prepare_upload(
        panel.registry,
        FOLDER / "ms_update1.xlsx",
        options={"export_date": "2026-10-01", "mode": "update"},
    )
    apply_preview(panel.registry, preview)
    store = new_calc.Store(panel.registry.rows)
    key = next(r for r in store.calc[("TP-M1", "1", 4)] if r["order"] == 6)
    assert store.plan(key["milestone_code"]) == "2027-01-31"  # дата новая без решения
    changed = new_calc.set_changes(store)
    assert changed and changed[0]["calc"] == "TP-M1/1/4"
    new = next(m["milestone_code"] for m in store.milestones if m["name"] == "Веха новая")
    assert changed[0]["add"] == [new] and not changed[0]["remove"]
    assert new not in {r["milestone_code"] for r in store.calc[("TP-M1", "1", 4)]}
    panel.save_calc_set_change(changed[0], "Протокол", "p.pdf")
    store = new_calc.Store(panel.registry.rows)
    assert new in {r["milestone_code"] for r in store.calc[("TP-M1", "1", 4)] if r["active"]}
    assert not new_calc.set_changes(store)


def test_c08_team_from_previous_stage_own_boundary(panel):
    create(panel)
    rows = [
        {
            "person_id": tab,
            "sov_type": "Вехи",
            "category": "РУПК",
            "role": "другая",
            "position": "Условная должность",
            "date_from": "2025-12-01",
            "source": "вручную",
        }
        for tab in ("T7001", "T7002")
    ]
    employees = [
        ("employee", {"person_id": t, "fio": f"Тестов {t}", "position": "Условная должность"})
        for t in ("T7001", "T7002")
    ]
    apply_records(panel.registry, employees, "сотрудники", "тест", "test")
    store = new_calc.Store(panel.registry.rows)
    for person in rows:
        records = []
        identity = new_calc.assignment_id(person["person_id"], "TP-M1", "1", 4)
        records.append(
            (
                "assignment",
                {
                    "id": identity,
                    "person_id": person["person_id"],
                    "project_code": "TP-M1",
                    "queue": "1",
                    "stage": 4,
                    "sov_type": "Вехи",
                    "category": "РУПК",
                    "role": "другая",
                    "position": "Условная должность",
                    "date_from": "2025-12-01",
                    "from_prev_stage": 0,
                },
            )
        )
        apply_records(panel.registry, records, "назначение", "тест", "test")
    key = next(m for m in store.milestones if m["name"] == "Получено РнВ")
    apply_records(
        panel.registry,
        [
            (
                "individual_date",
                {
                    "person_id": "T7002",
                    "milestone_code": key["milestone_code"],
                    "individual_plan": "2027-02-15",
                },
            )
        ],
        "инд. срок",
        "тест",
        "test",
    )
    draft = panel.new_calc_draft("TP-M1", "1", 5)
    assert draft["header"]["start_date"] == "2027-01-01"  # день после РнВ этапа 4
    team = {p["person_id"]: p for p in draft["team"]}
    assert team["T7001"]["date_from"] == "2027-01-01" and not team["T7001"]["individual"]
    assert team["T7002"]["date_from"] == "2027-02-16" and team["T7002"]["individual"]
    created = create(panel, stage=5, start=draft["header"]["start_date"], team=draft["team"])
    assert created["changed"] > 0
    new = [a for a in panel.registry.rows("assignment") if a["stage"] == 5]
    assert {(a["person_id"], a["date_from"], a["from_prev_stage"]) for a in new} == {
        ("T7001", "2027-01-01", 1),
        ("T7002", "2027-02-16", 1),
    }


def test_c09_creation_is_one_operation_and_rollback_removes_it(panel):
    before = {t: panel.registry.rows(t) for t in ("stage", "calc_milestone", "assignment")}
    result = create(panel)
    stage = panel.registry.rows("stage")
    assert len(stage) == 1 and stage[0]["key_milestone_code"]
    assert len(panel.registry.rows("calc_milestone")) == 6
    with pytest.raises(ValueError):
        create(panel)  # NEWCALC_EXISTS
    rollback_upload(panel.registry, result["upload_id"])
    assert {t: panel.registry.rows(t) for t in before} == before


def calc_data(panel, project="TP-M1", queue="1", stage=4, header=None):
    import json

    from sqlalchemy import select

    from sov.registry.models import milestone_raw_row

    store = new_calc.Store(panel.registry.rows)
    project_row = next(r for r in panel.registry.rows("project") if r["project_code"] == project)
    with panel.registry.engine.connect() as connection:
        raw = [
            (r.upload_id, r.row_no, json.loads(r.raw_json))
            for r in connection.execute(
                select(milestone_raw_row).order_by(milestone_raw_row.c.row_no)
            )
        ]
    return new_calc.calc_file_data(store, project_row, queue, stage, raw, header)


def variant(tmp_path, name, edit):
    """Условная копия шаблона с правкой XML (NEW_CALC §9, C10–C14)."""
    target = tmp_path / name
    with zipfile.ZipFile(TEMPLATE) as source, zipfile.ZipFile(target, "w") as output:
        for info in source.infolist():
            data = source.read(info.filename)
            output.writestr(info, edit(info.filename, data))
    return target


def test_c10_file_contents(panel, tmp_path):
    create(panel)
    template = variant(
        tmp_path,
        "без_пересчёта.xlsx",
        lambda n, d: d.replace(b' fullCalcOnLoad="1"', b"") if n == "xl/workbook.xml" else d,
    )
    data = calc_data(panel, header={"F1": True, "F2": "2025-12-01"})
    report = write_calculator_file(template, tmp_path / "out" / "bonus.xlsx", data)
    assert report["status"] == "создан" and not report["problems"]
    book = load_workbook(tmp_path / "out/bonus.xlsx")
    sheet = book["калькулятор"]
    assert (sheet["A1"].value, sheet["A2"].value, sheet["C16"].value) == (
        "Условный проект М1",
        "1",
        4,
    )
    middle = [sheet[f"E{n}"].value for n in range(16, 21)]
    assert middle == [
        "Веха А",
        "Общая веха проекта",
        "Веха Б",
        "Завершены работы ЖБ (часть 1)",
        "Завершены работы ЖБ (часть 2)",
    ]
    assert sheet["F17"].value is None  # веха проекта без очереди
    assert sheet["E36"].value == "Получено РнВ" and sheet["C36"].value == 4
    assert sheet["E21"].value is None  # прежние значения шаблона очищены
    hidden = {n for n, d in sheet.row_dimensions.items() if d.hidden}
    assert not hidden & set(range(16, 21))
    vehi = book["Лист для загрузки вех"]
    plans = [vehi.cell(r, 12).value for r in range(3, 18) if vehi.cell(r, 3).value]
    assert all(isinstance(p, str) and re.fullmatch(r"\d\d\.\d\d\.\d{4}", p) for p in plans if p)
    with zipfile.ZipFile(tmp_path / "out/bonus.xlsx") as result:
        assert b'fullCalcOnLoad="1"' in result.read("xl/workbook.xml")


def test_c11_only_three_parts_change(panel, tmp_path):
    create(panel)
    template = variant(
        tmp_path,
        "без_пересчёта.xlsx",
        lambda n, d: d.replace(b' fullCalcOnLoad="1"', b"") if n == "xl/workbook.xml" else d,
    )
    report = write_calculator_file(template, tmp_path / "bonus.xlsx", calc_data(panel))
    assert set(report["changed_parts"]) == {
        "xl/workbook.xml",
        "xl/worksheets/sheet1.xml",
        "xl/worksheets/sheet5.xml",
    }
    assert report["verified"] and verify(
        template, tmp_path / "bonus.xlsx", set(report["changed_parts"])
    )
    with zipfile.ZipFile(template) as a, zipfile.ZipFile(tmp_path / "bonus.xlsx") as b:
        same = [n for n in a.namelist() if n not in report["changed_parts"]]
        assert all(a.read(n) == b.read(n) for n in same)


def test_c12_key_row_by_label_not_color(panel, tmp_path):
    create(panel)
    data = calc_data(panel)
    # Другой цвет ключевой строки: меняем стиль ячеек строки 36.
    recolored = variant(
        tmp_path,
        "другой_цвет.xlsx",
        lambda n, d: (
            re.sub(rb'(<c r="[A-Z]+36") s="\d+"', rb'\1 s="1"', d)
            if n == "xl/worksheets/sheet1.xml"
            else d
        ),
    )
    assert write_calculator_file(recolored, tmp_path / "a.xlsx", data)["status"] == "создан"
    unlabeled = variant(
        tmp_path,
        "без_подписи.xlsx",
        lambda n, d: (
            re.sub(rb'<c r="M36"[^>]*?(?:/>|>.*?</c>)', b"", d, flags=re.S)
            if n == "xl/worksheets/sheet1.xml"
            else d
        ),
    )
    report = write_calculator_file(unlabeled, tmp_path / "b.xlsx", data)
    assert {p["code"] for p in report["problems"]} == {"XLS_KEY_ROW_NOT_FOUND"}
    assert not (tmp_path / "b.xlsx").exists()


def test_c13_overflow_and_short_key_formula_range(panel, tmp_path):
    create(panel)
    data = calc_data(panel)
    data["milestones"] = [{"name": f"Веха {i:02d}", "queue_raw": "1"} for i in range(25)]
    report = write_calculator_file(TEMPLATE, tmp_path / "a.xlsx", data)
    assert report["written"] == 21 and len(report["overflow"]) == 5  # 20 мест + ключевая
    assert "XLS_TOO_MANY_ROWS" in {p["code"] for p in report["problems"]}

    def short(name, data_bytes):
        if name != "xl/worksheets/sheet5.xml":
            return data_bytes
        # Формула ключа только в строках 3–5.
        return re.sub(
            rb'(<c r="A(\d+)"[^>]*?)>(<f[^>]*?(?:/>|>.*?</f>))',
            lambda m: m.group(0) if int(m.group(2)) <= 5 else m.group(1) + b">",
            data_bytes,
            flags=re.S,
        )

    report = write_calculator_file(
        variant(tmp_path, "короткий.xlsx", short), tmp_path / "b.xlsx", calc_data(panel)
    )
    misses = [p for p in report["problems"] if p["code"] == "XLS_LOOKUP_MISS"]
    assert misses and all(p["row"] for p in misses)


def test_c14_self_check_and_no_overwrite(panel, tmp_path):
    create(panel)
    data = calc_data(panel)
    report = write_calculator_file(TEMPLATE, tmp_path / "bonus.xlsx", data)
    assert len(report["check"]) == report["written"] == 6
    assert all(row["date"] for row in report["check"])
    stamp = (tmp_path / "bonus.xlsx").read_bytes()
    again = write_calculator_file(TEMPLATE, tmp_path / "bonus.xlsx", data)
    assert {p["code"] for p in again["problems"]} == {"XLS_FILE_EXISTS"}
    assert (tmp_path / "bonus.xlsx").read_bytes() == stamp


def test_export_through_service_logs(panel, tmp_path):
    create(panel)
    report = panel.export_calc_file(
        "TP-M1", "1", 4, TEMPLATE, tmp_path, "bonus_{project}_{queue}_{stage}"
    )
    assert report["status"] == "создан" and report["file"].endswith("bonus_TP-M1_1_4.xlsx")
    assert (
        TEMPLATE.read_bytes()
        == (ROOT / "reference/template/bonus_template_empty.xlsx").read_bytes()
    )


def test_calculation_uses_fixed_set_and_manual_after_accrued_change(panel):
    from sov.adapters.registry_case import RegistryAdapter

    create(panel)
    identity = new_calc.assignment_id("T7001", "TP-M1", "1", 4)
    apply_records(
        panel.registry,
        [
            (
                "employee",
                {"person_id": "T7001", "fio": "Тестов Семь", "position": "Условная должность"},
            ),
            (
                "assignment",
                {
                    "id": identity,
                    "person_id": "T7001",
                    "project_code": "TP-M1",
                    "queue": "1",
                    "stage": 4,
                    "sov_type": "Вехи",
                    "category": "РУПК",
                    "role": "другая",
                    "date_from": "2025-12-01",
                    "from_prev_stage": 0,
                },
            ),
        ],
        "команда",
        "тест",
        "test",
    )
    add_milestone(panel, "X-NEW", "TP-M1", "1", 4, "Веха вне набора", "2026-08-20")
    case = RegistryAdapter(panel.registry).build_case(identity)
    names_in_case = [m["name"] for m in case["milestones"]]
    assert "Веха вне набора" not in names_in_case  # набор зафиксирован при создании
    assert "Общая веха проекта" in names_in_case and not case["is_exception"]
    from sqlalchemy import insert

    from sov.registry.models import TABLES
    from sov.registry.session import timestamp

    with panel.registry.transaction() as connection:  # начисление закрытого квартала
        connection.execute(
            insert(TABLES["accrual"]).values(
                row_key="A1",
                id="A1",
                assignment_id=identity,
                kind="аванс за веху",
                valid_from=timestamp(),
                created_at=timestamp(),
            )
        )
    panel.save_calc_set_change(
        {"calc": "TP-M1/1/4", "add": ["X-NEW"], "remove": []}, "Протокол", "p.pdf"
    )
    case = RegistryAdapter(panel.registry).build_case(identity)
    assert "Веха вне набора" in [m["name"] for m in case["milestones"]]
    assert case["is_exception"]  # до правила переноса вех — ручной расчёт (NEW_CALC §4.3)
