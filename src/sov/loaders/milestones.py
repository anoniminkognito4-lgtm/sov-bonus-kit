"""Хранилище вех: выгрузка 1С → вехи реестра с версиями плана (docs/MILESTONES.md).

Чтение выгрузки (шапка в первых 10 строках, колонки по заголовкам), разбор этапов и
дат, натуральный ключ без кодов 1С, правила обновления §6.3 и проверки §7. Модуль
не пишет в базу: он готовит записи для общего механизма предпросмотра и применения.
"""

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

REQUIRED = (
    "Статус",
    "Наименование",
    "Ключевая",
    "Проект",
    "Очередь/корпус/строение",
    "Этап",
    "Плановый срок",
)
STORED = {"Шаблон": "template", "Вид": "kind", "Критерий закрытия": "criterion"}
KNOWN = {
    *REQUIRED,
    *STORED,
    "Фактический срок",
    "Дата закрытия в 1С",
    "Код проекта",
    "Код вехи",
}
PREFERRED_SHEET = "Лист для загрузки вех"
DONE, ASSIGNED = "Выполнена", "Назначена"
STAGE_STEMS = (("перв", 1), ("втор", 2), ("трет", 3), ("четверт", 4), ("пят", 5))
# Коды, при которых строка не применяется, пока пользователь не решит (§6.3, §7).
SKIP_CODES = {
    "MS_BAD_DATE",
    "MS_BAD_VALUE",
    "MS_DUPLICATE_KEY",
    "MS_PROJECT_UNMATCHED",
    "MS_FACT_CHANGED_CLOSED",
    "MS_FACT_REMOVED",
    "MS_KEY_CHANGED",
    "MS_POSSIBLY_RENAMED",
    "MS_POSSIBLY_MOVED",
}
BLOCKING = SKIP_CODES | {"MS_STATUS_UNKNOWN"}
MESSAGES = {
    "MS_BAD_DATE": "Дата не читается или не существует",
    "MS_BAD_VALUE": "Этап или признак «Ключевая» не распознан",
    "MS_NO_PLAN_DATE": "У вехи нет плановой даты",
    "MS_DUPLICATE_KEY": "Две строки с одним проектом, очередью, этапом и наименованием",
    "MS_PROJECT_UNMATCHED": "Проект или очередь не сопоставлены с реестром",
    "MS_NO_CODES": (
        "В выгрузке нет кодов вех: сопоставление по проекту, очереди, этапу и наименованию"
    ),
    "MS_NO_CLOSE_COLUMN": (
        "Нет колонки «Дата закрытия в 1С»: статус «Выполнена» считается закрытием, "
        "квартал начисления — по фактической дате"
    ),
    "MS_STATUS_UNKNOWN": "Статус 1С не «Назначена» и не «Выполнена»: по вехе не платим до решения",
    "MS_STATUS_FACT_MISMATCH": "«Выполнена» без факта или факт при статусе «Назначена»",
    "MS_STAGE_NO_KEY": "У этапа нет ключевой вехи",
    "MS_STAGE_MULTI_KEY": "В очереди и этапе несколько ключевых вех",
    "MS_PLAN_MOVED": "Перенос плановой даты",
    "MS_KEY_PLAN_MOVED": "Перенос плановой даты ключевой вехи: меняются длительность этапа и ПБ",
    "MS_FACT_CHANGED": "Фактическая дата изменилась",
    "MS_FACT_CHANGED_CLOSED": (
        "Факт изменился, а аванс по вехе уже в закрытом квартале: только корректировкой"
    ),
    "MS_FACT_REMOVED": "Факт пропал или статус вернулся в «Назначена»: не применяется",
    "MS_KEY_CHANGED": "Изменился признак «Ключевая»: не применяется",
    "MS_MISSING_IN_FILE": "Веха проекта из файла отсутствует в выгрузке",
    "MS_POSSIBLY_RENAMED": "Похоже на переименование вехи: нужно решение",
    "MS_POSSIBLY_MOVED": "Похоже на перенос вехи в другой этап: нужно решение",
    "MS_PART_PAIR": "Непарная часть вехи или вторая часть раньше первой",
    "MS_INDIVIDUAL_OUTDATED": "Индивидуальный срок раньше нового плана: действует ли он",
}


def clean(text):
    """trim, неразрывные пробелы и повторные пробелы — для заголовков и сравнения."""
    return re.sub(r"\s+", " ", str(text).replace("\xa0", " ")).strip() if text is not None else ""


def name_norm(text):
    return clean(text).lower().replace("ё", "е")


def parse_stage(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value == int(value) and 1 <= value <= 5:
            return int(value)
        raise ValueError("MS_BAD_VALUE")
    text = name_norm(value)
    match = re.fullmatch(r"(?:этап\s*)?([1-5])(?:\s*-?\s*(?:й|ый|ой|ий))?(?:\s*этап)?", text)
    if match:
        return int(match.group(1))
    for stem, number in STAGE_STEMS:
        if re.fullmatch(rf"{stem}\w*\s+этап", text):
            return number
    raise ValueError("MS_BAD_VALUE")


def parse_key_flag(value):
    text = name_norm(value)
    if text not in ("да", "нет"):
        raise ValueError("MS_BAD_VALUE")
    return 1 if text == "да" else 0


def parse_date(value):
    """Дата Excel, ISO или ДД.ММ.ГГГГ / ДД.ММ.ГГ с разделителями . / - \\ (§2.2)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not 1 <= value <= 2_958_465 or value != int(value):
            raise ValueError("MS_BAD_DATE")
        return (date(1899, 12, 30) + timedelta(days=int(value))).isoformat()
    text = clean(value)
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError:
        pass
    match = re.fullmatch(r"(\d{1,2})[./\\-](\d{1,2})[./\\-](\d{4}|\d{2})", text)
    if not match:
        raise ValueError("MS_BAD_DATE")
    day, month, year = (int(v) for v in match.groups())
    if len(match.group(3)) == 2:
        year += 2000 if year < 30 else 1900
    try:
        return date(year, month, day).isoformat()
    except ValueError as error:
        raise ValueError("MS_BAD_DATE") from error


def _header(rows):
    for index, row in enumerate(rows[:10]):
        cells = [clean(c) for c in row]
        if set(REQUIRED) <= set(cells):
            return index, cells
    return None


def read_workbook(path, sheet=None):
    """Найти лист выгрузки вех. None — это не выгрузка вех."""
    path = Path(path)
    if path.suffix.lower() in (".xls", ".xlsm", ".xlsb"):
        raise ValueError("Сохраните выгрузку вех как .xlsx")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        found = {}
        for worksheet in workbook.worksheets:
            rows = [tuple(r) for r in worksheet.iter_rows(values_only=True)]
            header = _header(rows)
            if header:
                found[worksheet.title] = (rows, header)
    finally:
        workbook.close()
    if not found:
        return None
    if sheet is None:
        sheet = PREFERRED_SHEET if PREFERRED_SHEET in found else next(iter(found))
    if sheet not in found:
        raise ValueError(f"На листе «{sheet}» нет шапки выгрузки вех")
    rows, (index, cells) = found[sheet]
    named = [c for c in cells if c]
    if len(named) != len(set(named)):
        raise ValueError("Повторяющиеся заголовки в шапке выгрузки вех")
    data = []
    for offset, values in enumerate(rows[index + 1 :], index + 2):
        raw = {cells[i]: v for i, v in enumerate(values) if i < len(cells) and cells[i]}
        if any(v is not None and str(v).strip() for v in raw.values()):
            data.append((offset, raw))
    layout = "MILESTONES_1C_CODES" if {"Код проекта", "Код вехи"} <= set(cells) else "MILESTONES_1C"
    return {
        "layout": layout,
        "sheet": sheet,
        "sheets": tuple(found),
        "header_row": index + 1,
        "headers": tuple(named),
        "rows": data,
    }


@dataclass
class Plan:
    records: list = field(default_factory=list)  # (строка, [(таблица, данные)])
    problems: list = field(default_factory=list)
    unmatched: list = field(default_factory=list)
    stats: dict = field(default_factory=dict)
    by_project: dict = field(default_factory=dict)


def _decisions(state, kind):
    rows = [r for r in state["manual_decision"].values() if r["kind"] == kind]
    return [json.loads(r["payload_json"]) for r in sorted(rows, key=lambda r: r["created_at"])]


def _latest_plan(state, code):
    versions = [r for r in state["milestone_plan_version"].values() if r["milestone_code"] == code]
    return max(versions, key=lambda r: r["version"]) if versions else None


def _ru(day):
    return date.fromisoformat(day).strftime("%d.%m.%Y")


def plan_upload(meta, state, file_hash, export_date, mode="update", sheet_name=None):
    """Подготовить записи и проблемы загрузки вех. `state` — таблицы {row_key: строка}."""
    plan = Plan()
    sheet = sheet_name or meta["sheet"]
    headers = set(meta["headers"])
    with_codes = meta["layout"] == "MILESTONES_1C_CODES"
    has_close = "Дата закрытия в 1С" in headers

    def problem(code, line=None, **extra):
        item = {
            "code": code,
            "blocking": code in BLOCKING,
            "message": MESSAGES[code],
            "source_ref": f"{sheet}!{line}" if line else None,
            "line": line,
        }
        item.update(extra)
        plan.problems.append(item)
        return item

    if not with_codes:
        problem("MS_NO_CODES")
    if not has_close:
        problem("MS_NO_CLOSE_COLUMN")
    project_alias = {name_norm(p["from"]): p["to"] for p in _decisions(state, "project_alias")}
    queue_alias = {
        (p["project_code"], name_norm(p["code_raw"])): p["queue"]
        for p in _decisions(state, "queue_alias")
    }
    status_alias = {p["raw"]: p["status"] for p in _decisions(state, "status_alias")}
    renames = _decisions(state, "milestone_rename")
    accepted_facts = {
        (p["milestone_code"], p.get("fact")) for p in _decisions(state, "fact_change_accept")
    }
    projects = list(state["project"].values())
    original = {r["milestone_code"]: r for r in state["milestone"].values()}
    keys = {r["natural_key"]: r["milestone_code"] for r in state["milestone_key"].values()}
    closed_advances = {
        r["milestone_code"] for r in state["accrual"].values() if r.get("milestone_code")
    }

    # 1. Разбор строк.
    parsed = []
    for line, raw in meta["rows"]:
        try:
            stage = parse_stage(raw.get("Этап"))
            is_key = parse_key_flag(raw.get("Ключевая"))
        except ValueError:
            problem("MS_BAD_VALUE", line)
            continue
        try:
            plan_date = parse_date(raw.get("Плановый срок"))
            fact = parse_date(raw.get("Фактический срок"))
            closed = parse_date(raw.get("Дата закрытия в 1С")) if has_close else None
        except ValueError:
            problem("MS_BAD_DATE", line)
            continue
        status = clean(raw.get("Статус"))
        name = clean(raw.get("Наименование"))
        queue_raw = clean(raw.get("Очередь/корпус/строение"))
        project_name = clean(raw.get("Проект"))
        code = clean(raw.get("Код проекта")) if with_codes else ""
        if not code:
            code = project_alias.get(name_norm(project_name))
        if not code:
            names = {
                p["project_code"]
                for p in projects
                if name_norm(p["name"]) == name_norm(project_name)
            }
            code = names.pop() if len(names) == 1 else None
        queues = {p["queue"] for p in projects if p["project_code"] == code}
        if not code or not queues:
            problem("MS_PROJECT_UNMATCHED", line, project=project_name)
            plan.unmatched.append(
                {"line": line, "row": raw, "message": MESSAGES["MS_PROJECT_UNMATCHED"]}
            )
            continue
        if not queue_raw:
            queue = ""  # веха проекта без очереди (§3.3)
        elif queue_raw in queues:
            queue = queue_raw
        else:
            queue = queue_alias.get((code, name_norm(queue_raw)))
            if queue not in queues:
                problem("MS_PROJECT_UNMATCHED", line, project=project_name, queue=queue_raw)
                plan.unmatched.append(
                    {"line": line, "row": raw, "message": "Очередь не сопоставлена с реестром"}
                )
                continue
        parsed.append(
            {
                "line": line,
                "raw": raw,
                "project_code": code,
                "queue": queue,
                "queue_raw": queue_raw,
                "stage": stage,
                "name": name,
                "is_key": is_key,
                "status": status,
                "plan": plan_date,
                "fact": fact,
                "closed": closed if has_close else (fact if status == DONE else None),
                "natural_key": json.dumps(
                    [code, name_norm(queue_raw), stage, name_norm(name)], ensure_ascii=False
                ),
                "file_code": clean(raw.get("Код вехи")) if with_codes else None,
            }
        )

    # 2. Дубли ключа в файле блокируют обе строки.
    seen = {}
    for row in parsed:
        seen.setdefault(row["file_code"] or row["natural_key"], []).append(row)
    duplicates = {k for k, rows in seen.items() if len(rows) > 1}
    for key in duplicates:
        for row in seen[key]:
            problem("MS_DUPLICATE_KEY", row["line"], milestone=row["name"])
    parsed = [r for r in parsed if (r["file_code"] or r["natural_key"]) not in duplicates]

    # 3. Код вехи: из файла, по натуральному ключу, по решению о переименовании, новый.
    counters = {}

    def new_code(row):
        prefix = f"{row['project_code']}/{row['queue'] or '-'}/{row['stage']}/"
        if prefix not in counters:
            used = [
                int(c[len(prefix) :])
                for c in original
                if c.startswith(prefix) and c[len(prefix) :].isdigit()
            ]
            counters[prefix] = max(used, default=0)
        counters[prefix] += 1
        return f"{prefix}{counters[prefix]:04d}"

    by_name = {}
    for m in original.values():
        by_name.setdefault(
            (m["project_code"], m["queue"], m["stage"], name_norm(m["name"])), []
        ).append(m["milestone_code"])
    for row in parsed:
        code = row["file_code"] or keys.get(row["natural_key"])
        if not code and not with_codes:
            # Веха уже в хранилище (например, загружена с кодами): связать ключ с ней.
            same = by_name.get(
                (row["project_code"], row["queue"], row["stage"], name_norm(row["name"]))
            )
            code = same[0] if same and len(same) == 1 else None
        if not code:
            for item in renames:
                target = original.get(item["milestone_code"])
                if (
                    target
                    and name_norm(item["new_name"]) == name_norm(row["name"])
                    and (target["project_code"], target["queue"], target["stage"])
                    == (row["project_code"], row["queue"], row["stage"])
                ):
                    code = target["milestone_code"]
        row["code"] = code
        row["new"] = code is None or code not in original
    file_codes = {r["code"] for r in parsed if r["code"]}
    file_projects = {r["project_code"] for r in parsed}
    missing = [
        m
        for m in original.values()
        if m["project_code"] in file_projects and m["milestone_code"] not in file_codes
    ]

    # 4. Новые строки: похоже на переименование или перенос в другой этап?
    paired = set()
    for row in parsed:
        if not row["new"]:
            continue
        if with_codes:
            continue  # коды 1С однозначны: новая веха — это новая веха
        same_stage = [
            m
            for m in missing
            if (m["project_code"], m["queue"], m["stage"])
            == (row["project_code"], row["queue"], row["stage"])
        ]
        renamed = [
            m
            for m in same_stage
            if (
                row["plan"]
                and (_latest_plan(state, m["milestone_code"]) or {}).get("plan_date") == row["plan"]
            )
            or (m.get("template") and m["template"] == clean(row["raw"].get("Шаблон")))
        ]
        moved = [
            m
            for m in missing
            if m["stage"] != row["stage"]
            and m["queue"] == row["queue"]
            and m["project_code"] == row["project_code"]
            and name_norm(m["name"]) == name_norm(row["name"])
        ]
        if renamed:
            row["skip"] = True
            paired.update(m["milestone_code"] for m in renamed)
            problem(
                "MS_POSSIBLY_RENAMED",
                row["line"],
                milestone=row["name"],
                proposal=f"переименование {renamed[0]['milestone_code']} или новая веха",
                milestone_code=renamed[0]["milestone_code"],
            )
        elif moved:
            row["skip"] = True
            paired.update(m["milestone_code"] for m in moved)
            problem(
                "MS_POSSIBLY_MOVED",
                row["line"],
                milestone=row["name"],
                proposal=(
                    f"веха {moved[0]['milestone_code']} перенесена из этапа {moved[0]['stage']}"
                ),
                milestone_code=moved[0]["milestone_code"],
            )
        else:
            row["code"] = row["code"] or new_code(row)
    for row in parsed:
        if not row.get("skip") and not row["code"]:
            row["code"] = new_code(row)
    for m in missing:
        if m["milestone_code"] not in paired:
            problem(
                "MS_MISSING_IN_FILE",
                milestone=m["name"],
                milestone_code=m["milestone_code"],
                project_code=m["project_code"],
            )

    # 5. Сравнение с хранилищем и записи.
    basis = (
        f"начальная загрузка, выгрузка 1С от {_ru(export_date)}"
        if mode == "seed"
        else f"выгрузка 1С от {_ru(export_date)}"
    )
    assignments = list(state["assignment"].values())
    stats_project = plan.by_project
    for row in parsed:
        project_stats = stats_project.setdefault(
            row["project_code"], {"added": 0, "moved": 0, "closed": 0, "missing": 0, "blocked": 0}
        )
        if row.get("skip"):
            project_stats["blocked"] += 1
            continue
        code, line = row["code"], row["line"]
        before = original.get(code)
        latest = _latest_plan(state, code)
        status_known = row["status"] in (DONE, ASSIGNED) or status_alias.get(row["status"]) in (
            "assigned",
            "done",
            "inactive",
        )
        if row["status"] and not status_known:
            problem("MS_STATUS_UNKNOWN", line, milestone=row["name"], status=row["status"])
        if (row["status"] == DONE and not row["fact"]) or (
            row["status"] == ASSIGNED and row["fact"]
        ):
            problem("MS_STATUS_FACT_MISMATCH", line, milestone=row["name"])
        if before:
            if before["is_key"] != row["is_key"]:
                problem("MS_KEY_CHANGED", line, milestone=row["name"], milestone_code=code)
                project_stats["blocked"] += 1
                continue
            removed = (before["fact_date"] and not row["fact"]) or (
                before["status"] == DONE and row["status"] == ASSIGNED
            )
            if removed and (code, row["fact"]) not in accepted_facts:
                problem("MS_FACT_REMOVED", line, milestone=row["name"], milestone_code=code)
                project_stats["blocked"] += 1
                continue
            if before["fact_date"] and row["fact"] and before["fact_date"] != row["fact"]:
                accepted = (code, row["fact"]) in accepted_facts
                if code in closed_advances and not accepted:
                    problem(
                        "MS_FACT_CHANGED_CLOSED",
                        line,
                        milestone=row["name"],
                        milestone_code=code,
                        before=before["fact_date"],
                        after=row["fact"],
                    )
                    project_stats["blocked"] += 1
                    continue
                problem(
                    "MS_FACT_CHANGED",
                    line,
                    milestone=row["name"],
                    milestone_code=code,
                    before=before["fact_date"],
                    after=row["fact"],
                )
        if not before:
            project_stats["added"] += 1
        if row["fact"] and not (before and before["fact_date"]):
            project_stats["closed"] += 1
        data = {
            "milestone_code": code,
            "project_code": row["project_code"],
            "queue": row["queue"],
            "stage": row["stage"],
            "name": row["name"],
            "is_key": row["is_key"],
            "status": row["status"],
            "fact_date": row["fact"],
            "closed_1c_date": row["closed"],
            "queue_raw": row["queue_raw"],
            "name_raw": str(row["raw"].get("Наименование")),
            **{column: clean(row["raw"].get(header)) or None for header, column in STORED.items()},
            "extra_json": json.dumps(
                {
                    k: (v.isoformat() if isinstance(v, (date, datetime)) else v)
                    for k, v in sorted(row["raw"].items())
                    if k not in KNOWN and v is not None
                },
                ensure_ascii=False,
                sort_keys=True,
            ),
            "first_file_hash": before["first_file_hash"]
            if before and before.get("first_file_hash")
            else file_hash,
        }
        records = [("milestone", data)]
        if not with_codes:
            records.append(
                ("milestone_key", {"natural_key": row["natural_key"], "milestone_code": code})
            )
        if not row["plan"]:
            problem("MS_NO_PLAN_DATE", line, milestone=row["name"], milestone_code=code)
        elif latest is None or latest["plan_date"] != row["plan"]:
            records.append(
                (
                    "milestone_plan_version",
                    {
                        "milestone_code": code,
                        "version": latest["version"] + 1 if latest else 1,
                        "plan_date": row["plan"],
                        "basis": basis,
                        "effective_date": export_date,
                        "file_hash": file_hash,
                    },
                )
            )
            if latest:
                shift = _days(row["plan"], latest["plan_date"])
                project_stats["moved"] += 1
                if row["is_key"]:
                    affected = sum(
                        1
                        for a in assignments
                        if (a["project_code"], a["queue"], a["stage"])
                        == (row["project_code"], row["queue"], row["stage"])
                    )
                    problem(
                        "MS_KEY_PLAN_MOVED",
                        line,
                        milestone=row["name"],
                        milestone_code=code,
                        shift_days=shift,
                        assignments=affected,
                    )
                else:
                    problem(
                        "MS_PLAN_MOVED",
                        line,
                        milestone=row["name"],
                        milestone_code=code,
                        shift_days=shift,
                    )
                for individual in state["individual_date"].values():
                    if (
                        individual["milestone_code"] == code
                        and individual["individual_plan"] < row["plan"]
                    ):
                        problem(
                            "MS_INDIVIDUAL_OUTDATED",
                            line,
                            milestone=row["name"],
                            milestone_code=code,
                            person_id=individual["person_id"],
                        )
        stage_row = next(
            (
                s
                for s in state["stage"].values()
                if (s["project_code"], s["queue"], s["stage"])
                == (row["project_code"], row["queue"], row["stage"])
            ),
            None,
        )
        if row["is_key"] and stage_row and stage_row.get("key_milestone_code") != code:
            records.append(("stage", {**stage_row, "key_milestone_code": code}))
        plan.records.append((line, records))
        row["code"] = code
    for m in missing:
        stats_project.setdefault(
            m["project_code"], {"added": 0, "moved": 0, "closed": 0, "missing": 0, "blocked": 0}
        )["missing"] += 1

    # 6. Проверки этапов и частей вех по итоговому набору проектов файла.
    result = {code: m for code, m in original.items() if m["project_code"] in file_projects}
    for _, records in plan.records:
        for table, data in records:
            if table == "milestone":
                result[data["milestone_code"]] = data
    stages = {}
    for m in result.values():
        stages.setdefault((m["project_code"], m["queue"], m["stage"]), []).append(m)
    plans = {code: (_latest_plan(state, code) or {}).get("plan_date") for code in result}
    for _, records in plan.records:
        for table, data in records:
            if table == "milestone_plan_version":
                plans[data["milestone_code"]] = data["plan_date"]
    for (project_code, queue, stage), rows in sorted(stages.items()):
        if queue:
            count = sum(1 for m in rows if m["is_key"])
            if count == 0:
                problem("MS_STAGE_NO_KEY", project_code=project_code, queue=queue, stage=stage)
            elif count > 1:
                problem("MS_STAGE_MULTI_KEY", project_code=project_code, queue=queue, stage=stage)
        parts = {}
        for m in rows:
            match = re.search(r"\(\s*часть\s*([12])\s*\)\s*$", m["name"], re.IGNORECASE)
            if match:
                base = name_norm(m["name"][: match.start()])
                parts.setdefault(base, {})[int(match.group(1))] = m
        for pair in parts.values():
            one, two = pair.get(1), pair.get(2)
            if not (one and two):
                problem("MS_PART_PAIR", milestone=(one or two)["name"], project_code=project_code)
            elif (
                plans.get(two["milestone_code"])
                and plans.get(one["milestone_code"])
                and (plans[two["milestone_code"]] < plans[one["milestone_code"]])
            ):
                problem("MS_PART_PAIR", milestone=two["name"], project_code=project_code)
    file_rows = [r for r in parsed]
    plan.stats = {
        "projects": len({r["project_code"] for r in file_rows}),
        "queues": len({(r["project_code"], r["queue"]) for r in file_rows if r["queue"]}),
        "stages": len({(r["project_code"], r["queue"], r["stage"]) for r in file_rows}),
        "milestones": len(file_rows),
        "key_milestones": sum(r["is_key"] for r in file_rows),
        "mode": mode,
        "export_date": export_date,
        "layout": meta["layout"],
        "sheet": sheet,
        "sheets": list(meta["sheets"]),
    }
    return plan


def _days(later, earlier):
    return (date.fromisoformat(later) - date.fromisoformat(earlier)).days


def main(argv=None):
    """uv run python -m sov.loaders.milestones seed <файл.xlsx> --export-date ДД.ММ.ГГГГ"""
    parser = argparse.ArgumentParser(description="Загрузка вех из выгрузки 1С")
    parser.add_argument("mode", choices=("seed", "update"))
    parser.add_argument("file")
    parser.add_argument("--export-date", required=True, help="ДД.ММ.ГГГГ")
    parser.add_argument("--db", default="build/sov.sqlite")
    parser.add_argument("--apply", action="store_true", help="применить после предпросмотра")
    args = parser.parse_args(argv)
    from sov.loaders.service import apply_preview, prepare_upload
    from sov.registry.session import Registry

    registry = Registry(args.db)
    options = {"export_date": parse_date(args.export_date), "mode": args.mode}
    preview = prepare_upload(registry, args.file, "командная строка", options)
    print(
        json.dumps(
            {"summary": preview.summary, "stats": dict(preview.stats)}, ensure_ascii=False, indent=1
        )
    )
    for item in preview.problems:
        print(f"{item['code']:24} {item.get('source_ref') or '':28} {item.get('milestone') or ''}")
    if args.apply:
        print(json.dumps(apply_preview(registry, preview), ensure_ascii=False))
    registry.dispose()
    return 0


if __name__ == "__main__":
    sys.exit(main())
