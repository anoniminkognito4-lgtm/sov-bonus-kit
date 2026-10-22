"""Выгрузка 1С УП ШР «Проектные команды» → назначения реестра (D-103).

Шапка: Проект | Очередь/строение | Мотивационный этап | ФИО | Подразделение | Должность |
Дата начала | Дата окончания | Покинул проект | Уволен | Категория | Общая расчетная загрузка.

- Проект и очередь — справочники project_alias и queue_alias (как для вех и расстановок).
- Этап словами («Второй этап» → 2) — разбор как у вех.
- Человек — общая функция ФИО → person_id (D-101); нового заводим, похожего — спрашиваем.
- Категория, вид СОВ и роль — по справочнику должностей; расхождение категории — вопрос.
- «Дата окончания» — плановый срок этапа, не дата ухода: сверяется с ключевой вехой.
- «Общая расчетная загрузка» не берётся: загрузка только по правилам СОВ.
- Даты ухода в файле нет: «Покинул проект» и «Уволен» сохраняются в назначении,
  а вопросы о датах остаются в проблемах расчёта, пока на них не ответят.

Модуль не пишет в базу: он готовит записи и проблемы для общего предпросмотра.
"""

import csv
import json
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

from sov.adapters.load import department_block
from sov.adapters.positions import PositionBook, resolve, role_for
from sov.engine.params import DEFAULT_PARAMS
from sov.loaders import persons
from sov.loaders.mapping import Unmatched, assignment_id, resolve_person, row_key
from sov.loaders.milestones import clean, name_norm, parse_date, parse_stage

HEADERS = (
    "Проект",
    "Очередь/строение",
    "Мотивационный этап",
    "ФИО",
    "Подразделение",
    "Должность",
    "Дата начала",
    "Дата окончания",
    "Покинул проект",
    "Уволен",
    "Категория",
    "Общая расчетная загрузка",
)
REQUIRED = HEADERS[:11]
CATEGORIES = ("КУПК (РП/ЗРП)", "КУПК", "РУПК", "МП")
# Строка с такой проблемой не применяется, пока её не разберут.
SKIP_CODES = {"UPSHR_BAD_VALUE", "UPSHR_PROJECT_UNMATCHED", "UPSHR_START_AFTER_END"}
MESSAGES = {
    "UPSHR_BAD_VALUE": "Значение не читается: этап, дата, «Да/Нет» или категория",
    "UPSHR_PROJECT_UNMATCHED": "Проект или очередь не сопоставлены с реестром",
    "UPSHR_START_AFTER_END": "Дата начала позже даты окончания",
    "UPSHR_DUPLICATE": "Повтор строки человека на том же этапе очереди: взята последняя",
    "UPSHR_FIO_SPACES": (
        "Лишние пробелы в ФИО: в макрокниге это был бы второй сотрудник, здесь — тот же"
    ),
    "UPSHR_LOAD_IGNORED": (
        "В колонке загрузки есть значение: не берётся, загрузка только по правилам СОВ"
        " (число больше 1 макрос «Данных» делил бы на 100)"
    ),
    "UPSHR_NO_END_DATE": ("Нет даты окончания: в макрокниге такой ключ пропускается целиком"),
    "UPSHR_NONSTANDARD_TERM": (
        "Дата окончания не совпадает с плановой датой ключевой вехи этапа: "
        "индивидуальный срок или опечатка?"
    ),
    "UPSHR_CATEGORY_MISMATCH": "Категория в 1С не совпадает со справочником должностей",
    "POSITION_UNRECOGNIZED": (
        "Должность не распознана по справочнику: нужен ответ "
        "(добавьте должность или исправление в справочник)"
    ),
    "LEAVE_DATE_MISSING": "В 1С «Покинул проект», а даты ухода нет: укажите дату ухода",
    "LEAVE_ACTIVE_IN_1C": "В 1С человек ещё на проекте, а в реестре есть дата ухода",
    "DISMISSAL_DATE_MISSING": "В 1С «Уволен», а даты увольнения из кадров нет: укажите её",
    **persons.QUESTIONS,
}
BLOCKING = SKIP_CODES | set(persons.QUESTIONS) | {"LEAVE_DATE_MISSING", "DISMISSAL_DATE_MISSING"}


def _find_header(rows):
    for index, row in enumerate(rows[:10]):
        cells = [clean(c) for c in row]
        if set(REQUIRED) <= set(cells):
            return index, cells
    return None


def read(path):
    """Найти шапку УП ШР в первых 10 строках csv или xlsx. None — это не УП ШР."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = [tuple(r) for r in csv.reader(stream, delimiter=";")]
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            rows = []
            for worksheet in workbook.worksheets:
                rows = [tuple(r) for r in worksheet.iter_rows(values_only=True)]
                if _find_header(rows):
                    break
        finally:
            workbook.close()
    else:
        return None
    header = _find_header(rows)
    if not header:
        return None
    index, cells = header
    data = []
    for offset, values in enumerate(rows[index + 1 :], index + 2):
        raw = {cells[i]: v for i, v in enumerate(values) if i < len(cells) and cells[i]}
        if any(v is not None and str(v).strip() for v in raw.values()):
            data.append((offset, raw))
    return {"header_row": index + 1, "rows": data}


@dataclass
class Plan:
    records: list = field(default_factory=list)  # (строка, [(таблица, данные)])
    problems: list = field(default_factory=list)
    unmatched: list = field(default_factory=list)
    stats: dict = field(default_factory=dict)


def _decisions(state, kind):
    rows = [r for r in state["manual_decision"].values() if r["kind"] == kind]
    return [json.loads(r["payload_json"]) for r in sorted(rows, key=lambda r: r["created_at"])]


def _yes_no(value):
    text = name_norm(value)
    if text in ("да", "1", "true"):
        return 1
    if text in ("нет", "0", "false", ""):
        return 0
    raise ValueError("UPSHR_BAD_VALUE")


def _day(value, shift):
    return (date.fromisoformat(value) + timedelta(days=shift)).isoformat()


def _key_plan(state, project, queue, stage, person_id):
    """Плановая дата ключевой вехи этапа; индивидуальный срок человека, если он есть."""
    rows = [
        s
        for s in state["stage"].values()
        if (s["project_code"], s["queue"], s["stage"]) == (project, queue, stage)
    ]
    code = rows[0]["key_milestone_code"] if rows else None
    if not code:
        keys = [
            m
            for m in state["milestone"].values()
            if (m["project_code"], m["queue"], m["stage"]) == (project, queue, stage)
            and m["is_key"]
        ]
        code = keys[0]["milestone_code"] if len(keys) == 1 else None
    if not code:
        return None, None
    own = [
        r["individual_plan"]
        for r in state["individual_date"].values()
        if r["person_id"] == person_id and r["milestone_code"] == code
    ]
    versions = [r for r in state["milestone_plan_version"].values() if r["milestone_code"] == code]
    plan = max(versions, key=lambda r: r["version"])["plan_date"] if versions else None
    return plan, (own[0] if own else None)


def plan_upload(meta, state, params=DEFAULT_PARAMS):
    """Записи и проблемы загрузки УП ШР. `state` — таблицы {row_key: строка}."""
    plan = Plan()

    def problem(code, line=None, **extra):
        item = {
            "code": code,
            "blocking": code in BLOCKING,
            "message": MESSAGES[code],
            "line": line,
            "source_ref": f"строка {line}" if line else None,
        }
        item.update(extra)
        plan.problems.append(item)
        return item

    project_alias = {name_norm(p["from"]): p["to"] for p in _decisions(state, "project_alias")}
    queue_alias = {
        (p["project_code"], name_norm(p["code_raw"])): p["queue"]
        for p in _decisions(state, "queue_alias")
    }
    projects = list(state["project"].values())
    book = PositionBook(
        list(state["position_ref"].values()),
        list(state["position_fix"].values()),
        [p["name"] for p in projects] + [p["project_code"] for p in projects],
    )
    new_people = {}  # нормализованное ФИО → (person_id, строка сотрудника) новых в этом файле
    seen = {}
    rows = []
    for line, raw in meta["rows"]:
        try:
            stage = parse_stage(raw.get("Мотивационный этап"))
            start = parse_date(raw.get("Дата начала"))
            end = parse_date(raw.get("Дата окончания"))
            left, dismissed = _yes_no(raw.get("Покинул проект")), _yes_no(raw.get("Уволен"))
        except ValueError:
            problem("UPSHR_BAD_VALUE", line)
            plan.unmatched.append(
                {"line": line, "row": raw, "message": MESSAGES["UPSHR_BAD_VALUE"]}
            )
            continue
        category = clean(raw.get("Категория")) or None
        if not start or category not in CATEGORIES:
            problem("UPSHR_BAD_VALUE", line)
            plan.unmatched.append(
                {"line": line, "row": raw, "message": MESSAGES["UPSHR_BAD_VALUE"]}
            )
            continue
        if end and start > end:
            problem("UPSHR_START_AFTER_END", line)
            plan.unmatched.append(
                {"line": line, "row": raw, "message": MESSAGES["UPSHR_START_AFTER_END"]}
            )
            continue

        # Проект и очередь: справочники соответствий, затем точное название.
        project_name, queue_raw = clean(raw.get("Проект")), clean(raw.get("Очередь/строение"))
        code = project_alias.get(name_norm(project_name))
        if not code:
            names = {
                p["project_code"]
                for p in projects
                if name_norm(p["name"]) == name_norm(project_name)
                or p["project_code"] == project_name
            }
            code = names.pop() if len(names) == 1 else None
        queues = sorted({p["queue"] for p in projects if p["project_code"] == code})
        if not queue_raw and len(queues) == 1:
            queue = queues[0]  # у проекта одна очередь — выбирать не из чего
        elif queue_raw in queues:
            queue = queue_raw
        else:
            queue = queue_alias.get((code, name_norm(queue_raw)))
        if not code or queue not in queues:
            problem("UPSHR_PROJECT_UNMATCHED", line, project=project_name, queue=queue_raw)
            plan.unmatched.append(
                {"line": line, "row": raw, "message": MESSAGES["UPSHR_PROJECT_UNMATCHED"]}
            )
            continue

        # Человек.
        fio_raw = raw.get("ФИО")
        fio = clean(fio_raw)
        department, position_raw = clean(raw.get("Подразделение")), clean(raw.get("Должность"))
        if fio and str(fio_raw) != fio:
            problem("UPSHR_FIO_SPACES", line, fio=fio)
        if persons.norm(fio) in new_people:
            person_id, created = new_people[persons.norm(fio)][0], False
        else:
            try:
                person_id, created = resolve_person(
                    state,
                    fio,
                    department,
                    position_raw,
                    create=True,
                    reserved=[p for p, _ in new_people.values()],
                )
            except Unmatched as error:
                if error.problem:
                    plan.problems.append({**error.problem, "line": line})
                plan.unmatched.append({"line": line, "row": raw, "message": str(error)})
                continue
            except ValueError as error:
                plan.unmatched.append({"line": line, "row": raw, "message": str(error)})
                continue
        if created:
            employee = {
                "person_id": person_id,
                "fio": fio,
                "department": department or None,
                "position": position_raw or None,
            }
            new_people[persons.norm(fio)] = (person_id, employee)
            plan.records.append((line, [("employee", employee)]))
        employee = state["employee"].get(row_key("employee", {"person_id": person_id}))
        employee = employee or new_people[persons.norm(fio)][1]

        # Должность: категория, вид СОВ и роль по справочнику.
        block = department_block(department, params)
        found = book.recognize(position_raw, block) if block else {"ref": None}
        ref = found["ref"]
        if book and block and ref is None:
            problem("POSITION_UNRECOGNIZED", line, position=position_raw, block=block)
        if ref and ref.get("category") and category and ref["category"] != category:
            problem(
                "UPSHR_CATEGORY_MISMATCH",
                line,
                fio=fio,
                position=ref["position"],
                before=ref["category"],
                after=category,
            )
        filled = resolve({"category": category, "sov_type": None}, position_raw, ref)

        if raw.get("Общая расчетная загрузка") not in (None, ""):
            problem("UPSHR_LOAD_IGNORED", line, value=raw.get("Общая расчетная загрузка"))
        if not end:
            problem("UPSHR_NO_END_DATE", line, fio=fio)
        else:
            key_plan, individual = _key_plan(state, code, queue, stage, person_id)
            if key_plan and end != key_plan and end != individual:
                problem(
                    "UPSHR_NONSTANDARD_TERM",
                    line,
                    fio=fio,
                    project_code=code,
                    before=key_plan,
                    after=end,
                    proposal="индивидуальный срок по ключевой вехе или исправить в 1С",
                )
        identity = assignment_id(person_id, code, queue, stage)
        rows.append(
            {
                "line": line,
                "raw": raw,
                "identity": identity,
                "person_id": person_id,
                "employee": employee,
                "project_code": code,
                "queue": queue,
                "stage": stage,
                "start": start,
                "end": end,
                "left": left,
                "dismissed": dismissed,
                "category": filled["category"],
                "sov_type": filled["sov_type"],
                "role": role_for(ref),
                "position": filled["position"] or None,
                "fio": fio,
            }
        )
        if identity in seen:
            problem("UPSHR_DUPLICATE", line, fio=fio)
        seen[identity] = rows[-1]

    # Переход с этапа N−1: назначение там закончилось за день до «Даты начала».
    by_identity = {r["identity"]: r for r in rows}
    for row in seen.values():
        previous_id = assignment_id(
            row["person_id"], row["project_code"], row["queue"], row["stage"] - 1
        )
        previous = by_identity.get(previous_id) or next(
            (a for a in state["assignment"].values() if a["id"] == previous_id), None
        )
        previous_end = None
        if previous is not None:
            previous_end = previous.get("end") or previous.get("date_to")
            if previous_end is None:
                previous_end = _key_plan(
                    state, row["project_code"], row["queue"], row["stage"] - 1, row["person_id"]
                )[0]
        from_prev = int(bool(previous_end) and _day(previous_end, 1) == row["start"])

        before = next((a for a in state["assignment"].values() if a["id"] == row["identity"]), None)
        date_to, end_reason = row["end"], None
        if before and before.get("end_reason"):
            # Ручное решение (или расстановка) главнее файла: дата ухода не меняется.
            date_to, end_reason = before["date_to"], before["end_reason"]
            if not row["left"] and not row["dismissed"]:
                problem(
                    "LEAVE_ACTIVE_IN_1C",
                    row["line"],
                    fio=row["fio"],
                    before=before["date_to"],
                    after="в 1С: на проекте",
                )
        elif row["left"]:
            problem(
                "LEAVE_DATE_MISSING", row["line"], fio=row["fio"], project_code=row["project_code"]
            )
        if row["dismissed"] and not row["employee"].get("dismissal_date"):
            problem("DISMISSAL_DATE_MISSING", row["line"], fio=row["fio"])
        assignment = {
            "id": row["identity"],
            "person_id": row["person_id"],
            "project_code": row["project_code"],
            "queue": row["queue"],
            "stage": row["stage"],
            "sov_type": row["sov_type"],
            "category": row["category"],
            "role": row["role"],
            "position": row["position"],
            "date_from": row["start"],
            "date_to": date_to,
            "end_reason": end_reason,
            "from_prev_stage": from_prev,
            "left_1c": row["left"],
            "dismissed_1c": row["dismissed"],
        }
        plan.records.append((row["line"], [("assignment", assignment)]))
    plan.stats = {
        "rows": len(meta["rows"]),
        "assignments": len(seen),
        "new_people": len(new_people),
        "questions": sum(p["blocking"] for p in plan.problems),
    }
    return plan
