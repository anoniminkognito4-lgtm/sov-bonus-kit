"""Распознавание по заголовкам, приведение дат/чисел и чтение CSV/xlsx."""

import csv
import hashlib
import math
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

FORMATS = {
    "sov_param": "edition name value",
    "employee": (
        "tab_no fio position department hire_date dismissal_date probation_end probation_result "
    ),
    "bgb_period": "tab_no bgb date_from date_to",
    "project": "project_code name queue blocks is_exception sale_date handover_to_completion_date",
    "stage": "project_code queue stage start_date key_milestone is_completion",
    "assignment": (
        "tab_no project_code queue stage sov_type category role position_on_project "
        "date_from date_to end_reason load from_prev_stage "
    ),
    "individual_date": "tab_no project_code stage milestone individual_plan basis",
    "npv": "project_code quarter fm_version npv_plan_paur npv_forecast_paur fub_approved_date",
    "start_npv": "tab_no project_code stage npv_start fixed_date basis",
    "bs_share": "project_code quarter share_transferred",
    "kpi": "tab_no quarter kpi",
    "stage5_expense": "project_code queue stage expense_remaining plan_budget fact_budget",
    "key_rate": "date_from rate",
    "holiday": "date kind",
    "redistribution": "tab_no project_code stage milestone pct decided_by approved_date basis",
    "rp_project_set": (
        "tab_no rank project_code queue class area_k_m2 stage extra_objects is_queue "
        "is_logistics approved_date approved_by "
    ),
    "position_fix": "wrong right",
    "case_map": "case_id tab_no project_code",
    "payroll_example": "Табельный номер ФИО Подразделение",
}
MS_HEADERS = [
    "Статус",
    "Наименование",
    "Шаблон",
    "Ключевая",
    "Вид",
    "Проект",
    "Очередь/корпус/строение",
    "Этап",
    "Заказчик",
    "Исполнитель",
    "Плановый срок",
    "Критерий закрытия",
    "Фактический срок",
]
# Справочник должностей (samples/positions.csv): заголовки по-русски.
POSITION_HEADERS = ("Подразделение", "Должность", "Категория", "Вид СОВ по умолчанию")
DATE_FIELDS = {
    "hire_date",
    "dismissal_date",
    "probation_end",
    "sale_date",
    "handover_to_completion_date",
    "start_date",
    "date_from",
    "date_to",
    "individual_plan",
    "fub_approved_date",
    "fixed_date",
    "date",
    "approved_date",
}
FLOAT_FIELDS = {
    "bgb",
    "load",
    "npv_plan_paur",
    "npv_forecast_paur",
    "npv_start",
    "share_transferred",
    "kpi",
    "plan_budget",
    "fact_budget",
    "rate",
    "pct",
    "area_k_m2",
}
INT_FIELDS = {"stage", "rank", "extra_objects"}
BOOL_FIELDS = {
    "is_exception",
    "is_completion",
    "from_prev_stage",
    "expense_remaining",
    "is_queue",
    "is_logistics",
}


def normalize_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    value = str(value).strip()
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError:
        return datetime.strptime(value, "%d.%m.%Y").date().isoformat()


def normalize_bool(value):
    values = {"да": 1, "нет": 0, "true": 1, "false": 0, "1": 1, "0": 0}
    if str(value).strip().lower() not in values:
        raise ValueError(f"Неверный логический признак: {value}")
    return values[str(value).strip().lower()]


def normalize_row(raw):
    row = {}
    for key, value in raw.items():
        value = value.strip() if isinstance(value, str) else value
        if value in ("", None):
            row[key] = None
        elif key in DATE_FIELDS:
            row[key] = normalize_date(value)
        elif key in BOOL_FIELDS:
            row[key] = normalize_bool(value)
        elif key in FLOAT_FIELDS | INT_FIELDS:
            number = float(str(value).replace(" ", "").replace(",", "."))
            if not math.isfinite(number):
                raise ValueError(f"Число {key} должно быть конечным")
            if key in INT_FIELDS and number != int(number):
                raise ValueError(f"{key} должен быть целым")
            row[key] = int(number) if key in INT_FIELDS else number
        else:
            row[key] = str(value)
    if row.get("stage") is not None and not 1 <= row["stage"] <= 5:
        raise ValueError("Этап должен быть от 1 до 5")
    if row.get("quarter") and not re.fullmatch(r"\d{4}-Q[1-4]", row["quarter"]):
        raise ValueError("Неверный формат квартала")
    if row.get("date_from") and row.get("date_to") and row["date_to"] < row["date_from"]:
        raise ValueError("Конец периода раньше начала")
    if row.get("probation_result"):
        values = {
            "пройден": "пройден",
            "не пройден": "не пройден",
            "проходит": "проходит",
            "passed": "пройден",
            "failed": "не пройден",
            "pending": "проходит",
        }
        value = values.get(row["probation_result"].lower())
        if value is None or not row.get("probation_end"):
            raise ValueError("Нужны дата окончания и результат ИС: пройден / не пройден / проходит")
        row["probation_result"] = value
    for field in ("share_transferred",):
        if row.get(field) is not None and not 0 <= row[field] <= 1:
            raise ValueError(f"{field} должен быть от 0 до 1")
    if row.get("pct") is not None and abs(row["pct"]) > 0.3:
        raise ValueError("Перераспределение не может превышать ±30%")
    for field in (
        "bgb",
        "load",
        "rate",
        "area_k_m2",
        "extra_objects",
        "fact_budget",
        "plan_budget",
    ):
        if row.get(field) is not None and row[field] < 0:
            raise ValueError(f"{field} не может быть отрицательным")
    return row


def read_source(path):
    path = Path(path)
    if path.suffix.lower() not in (".csv", ".xlsx"):
        raise ValueError("Поддерживаются только условные CSV и xlsx без макросов")
    file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream, delimiter=";"))
    else:
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            worksheet = (
                workbook["Лист для загрузки вех"]
                if "Лист для загрузки вех" in workbook.sheetnames
                else workbook.active
            )
            rows = list(worksheet.values)
        finally:
            workbook.close()
    if not rows:
        raise ValueError("Нет строки заголовков")
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    if len(headers) != len(set(headers)) or "" in headers:
        raise ValueError("Пустые или повторяющиеся заголовки")
    if set(MS_HEADERS) <= set(headers):
        kind = "milestone"
    elif set(POSITION_HEADERS) <= set(headers):
        kind = "position_ref"
    else:
        candidates = [
            (name, columns.split())
            for name, columns in FORMATS.items()
            if set(columns.split()) <= set(headers)
        ]
        if not candidates:
            raise ValueError("Тип не распознан: отсутствуют обязательные колонки")
        kind = max(candidates, key=lambda item: len(item[1]))[0]
    result, errors = [], []
    for line, values in enumerate(rows[1:], 2):
        if not any(v is not None and str(v).strip() for v in values):
            continue
        try:
            if len(values) != len(headers):
                raise ValueError("Количество значений не совпадает с заголовками")
            raw = dict(zip(headers, values, strict=True))
            result.append((line, raw if kind == "milestone" else normalize_row(raw)))
        except (ValueError, TypeError) as error:
            errors.append({"line": line, "message": str(error)})
    return kind, file_hash, result, errors
