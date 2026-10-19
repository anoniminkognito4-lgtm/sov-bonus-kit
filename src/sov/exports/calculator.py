"""Калькулятор по проекту и этапу: лист как в шаблоне, только значения (DATA_FORMATS §6)."""

from datetime import date

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from sov.exports.files import money, new_workbook, stable_bytes

DARK = "173F45"
LIGHT = "E8F0EE"
THIN = Side(style="thin", color="C9D6D3")

MILESTONE_COLUMNS = (
    "№",
    "Веха / часть",
    "Плановый срок",
    "Инд. срок",
    "Факт / событие",
    "Откл., дн",
    "Вес",
    "Сумма, ₽",
    "Проценты, ₽",
    "Дата выплаты",
    "Статус",
)
WIDTHS = (6, 34, 14, 14, 14, 10, 10, 16, 14, 14, 22)


def _value(value):
    if isinstance(value, str) and len(value) == 10 and value[4] == "-" and value[7] == "-":
        try:
            return date.fromisoformat(value)
        except ValueError:
            return value
    return value


def _rub(value):
    return float(money(value)) if value is not None else None


def _style_cells(sheet, row, count, fill=None, bold=False, color=None):
    for column in range(1, count + 1):
        cell = sheet.cell(row, column)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if bold or color:
            cell.font = Font(bold=bold, color=color)


def _formats(sheet, row):
    for column in range(1, len(MILESTONE_COLUMNS) + 1):
        cell = sheet.cell(row, column)
        if isinstance(cell.value, date):
            cell.number_format = "DD.MM.YYYY"
        elif isinstance(cell.value, float):
            cell.number_format = "0.00%" if column == 7 else "#,##0.00"


def _employee_block(sheet, row, value):
    case, employee = value["case"], value["result"]["employees"][0]
    source = case["employees"][0]
    title = f"{employee['fio']} · {employee['tab_no']} · {employee['category']}"
    if employee.get("role") and employee["role"] != "другая":
        title += f" · {employee['role']}"
    sheet.cell(row, 1, title)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(MILESTONE_COLUMNS))
    _style_cells(sheet, row, len(MILESTONE_COLUMNS), fill=DARK, bold=True, color="FFFFFF")
    row += 1
    facts = [
        ("Дата начала работы на этапе", _value(employee["start"])),
        ("Дата ухода", _value(employee.get("leave_date"))),
        ("Переход из предыдущего этапа", "да" if case.get("continues_from_prev_stage") else "нет"),
        ("Длительность этапа, дней", employee.get("duration_days")),
        ("Проектный бонус (ПБ) этапа, ₽", _rub(employee.get("pb"))),
        ("Сумма на авансы по вехам, ₽", _rub(employee.get("advance_pool"))),
        ("Остаток на финал этапа, ₽", _rub(employee.get("final_pool"))),
        ("NPV стартовый", source.get("npv_start")),
        ("NPV текущий", source.get("npv_current")),
        ("К фин. результата", employee.get("k_fin")),
        ("Множитель БС", employee.get("k_bs")),
        ("Проценты на отложенную часть, ₽", _rub(employee["interest"]["total"])),
        ("Итого к выплате (рассчитано), ₽", _rub(employee.get("total_payable"))),
    ]
    for label, item in facts:
        if item is None:
            continue
        sheet.cell(row, 2, label)
        cell = sheet.cell(row, 4, item)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        if isinstance(item, date):
            cell.number_format = "DD.MM.YYYY"
        elif isinstance(item, float):
            cell.number_format = "#,##0.00" if abs(item) >= 10 else "0.0000"
        row += 1
    row += 1
    for column, label in enumerate(MILESTONE_COLUMNS, 1):
        sheet.cell(row, column, label)
    _style_cells(sheet, row, len(MILESTONE_COLUMNS), fill=LIGHT, bold=True)
    row += 1
    for number, item in enumerate(employee["milestones"], 1):
        values = (
            number,
            item["name"],
            _value(item.get("plan")),
            _value(item.get("individual_plan")),
            _value(item.get("fact")),
            item.get("deviation_days"),
            item.get("weight"),
            _rub(item.get("amount")),
            None,
            _value(item.get("pay_date")),
            item["calc_status"],
        )
        for column, cell_value in enumerate(values, 1):
            sheet.cell(row, column, cell_value)
        _style_cells(sheet, row, len(MILESTONE_COLUMNS))
        _formats(sheet, row)
        row += 1
    final = employee["final"]
    values = (
        None,
        f"Ключевая веха · {final['kind']}",
        _value(final.get("plan")),
        _value(final.get("individual_plan")),
        _value(final.get("fact")),
        final.get("deviation_days"),
        None,
        _rub(final.get("amount")),
        None,
        _value(final.get("pay_date")),
        final["calc_status"],
    )
    for column, cell_value in enumerate(values, 1):
        sheet.cell(row, column, cell_value)
    _style_cells(sheet, row, len(MILESTONE_COLUMNS), bold=True)
    _formats(sheet, row)
    row += 1
    for part in employee["payouts"]:
        values = (
            None,
            f"Финал: часть {part['part']}",
            None,
            None,
            _value(part.get("event_date")),
            None,
            None,
            _rub(part.get("amount")),
            _rub(part.get("interest")),
            _value(part.get("pay_date")),
            part["calc_status"],
        )
        for column, cell_value in enumerate(values, 1):
            sheet.cell(row, column, cell_value)
        _style_cells(sheet, row, len(MILESTONE_COLUMNS))
        _formats(sheet, row)
        row += 1
    periods = source.get("periods") or []
    if periods:
        row += 1
        for column, label in enumerate(
            ("", "Период БГБ и загрузки", "С", "По", "", "", "Загрузка"), 1
        ):
            sheet.cell(row, column, label or None)
        sheet.cell(row, 8, "БГБ, ₽")
        _style_cells(sheet, row, 8, fill=LIGHT, bold=True)
        row += 1
        for period in periods:
            sheet.cell(row, 3, _value(period["from"])).number_format = "DD.MM.YYYY"
            sheet.cell(row, 4, _value(period["to"])).number_format = "DD.MM.YYYY"
            sheet.cell(row, 7, period["load"]).number_format = "0.00%"
            sheet.cell(row, 8, period["bgb"]).number_format = "#,##0.00"
            row += 1
    return row + 2


def calculator_workbook(results, selections, labels=None, quarter=None, preliminary=True):
    """`results` — результаты PanelService.calculate; `selections` — (код, очередь, этап)."""
    workbook = new_workbook()
    used = set()
    for code, queue, stage in selections:
        values = [
            v
            for v in results.values()
            if (v["case"]["project_code"], v["case"]["queue"], v["case"]["stage"])
            == (code, queue, stage)
        ]
        values.sort(key=lambda v: v["result"]["employees"][0]["fio"] or "")
        name = f"{code} оч{queue} эт{stage}"[:28]
        base, suffix = name, 2
        while name in used:
            name, suffix = f"{base}_{suffix}", suffix + 1
        used.add(name)
        sheet = workbook.create_sheet(name)
        label = (labels or {}).get(
            (code, queue, stage)
        ) or f"{code} · очередь {queue} · этап {stage}"
        sheet.cell(1, 1, f"Калькулятор · {label}").font = Font(bold=True, size=14, color=DARK)
        status = "Расчёт предварительный" if preliminary else "Расчёт на закрытие квартала"
        sheet.cell(2, 1, f"{status}{f' · квартал {quarter}' if quarter else ''}")
        sheet.cell(2, 1).font = Font(italic=True, color="5B6B68")
        row = 4
        if not values:
            sheet.cell(row, 1, "Нет назначений для расчёта")
        for value in values:
            row = _employee_block(sheet, row, value)
        for index, width in enumerate(WIDTHS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.sheet_view.showGridLines = False
    if not workbook.sheetnames:
        workbook.create_sheet("калькулятор").cell(1, 1, "Ничего не выбрано")
    return stable_bytes(workbook)
