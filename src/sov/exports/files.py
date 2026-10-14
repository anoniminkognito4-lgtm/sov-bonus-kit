"""Выгрузки значений без формул; округление денег только здесь."""

import csv
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PAYROLL_COLUMNS = (
    "Табельный номер",
    "ФИО",
    "Подразделение",
    "Проект, очередь",
    "Этап",
    "Вид начисления",
    "Веха",
    "Квартал начисления",
    "Дата выплаты",
    "Сумма, ₽",
    "ID расчёта",
)


def money(value):
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def safe_text(value):
    # Текст остаётся текстом при открытии CSV в Excel.
    return "'" + value if isinstance(value, str) and value[:1] in ("=", "+", "-", "@") else value


def csv_bytes(columns, rows):
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=";")
    writer.writerow(columns)
    for row in rows:
        writer.writerow(
            [
                str(value).replace(".", ",") if isinstance(value, Decimal) else safe_text(value)
                for value in row
            ]
        )
    return output.getvalue().encode("utf-8-sig")


def excel_bytes(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, (columns, rows) in sheets.items():
        sheet = workbook.create_sheet(name[:31])
        sheet.append(list(columns))
        for row in rows:
            sheet.append([float(value) if isinstance(value, Decimal) else value for value in row])
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173F45")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 32
        for index, column in enumerate(sheet.columns, 1):
            width = min(48, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[get_column_letter(index)].width = width
    result = BytesIO()
    workbook.save(result)
    workbook.close()
    return result.getvalue()


def payroll_rows(closed):
    rows = []
    for entry in closed["payload"]["entries"]:
        if entry["calc_status"] != "рассчитано" or not entry["amount"]:
            continue
        rows.append(
            [
                entry["tab_no"],
                entry["fio"],
                entry["department"],
                f"{entry['project']}, очередь {entry['queue']}",
                entry["stage"],
                entry["kind"],
                entry["milestone"],
                entry["quarter"],
                date.fromisoformat(entry["pay_date"]).strftime("%d.%m.%Y"),
                money(entry["amount"]),
                f"{closed['id']} / {entry['edition']}",
            ]
        )
    return rows


def payroll(closed, format="xlsx"):
    rows = payroll_rows(closed)
    if format == "csv":
        return csv_bytes(PAYROLL_COLUMNS, rows)
    return excel_bytes({"К начислению": (PAYROLL_COLUMNS, rows)})


def statement(results, project_code, stage, preliminary=True):
    columns = (
        "Табельный номер",
        "ФИО",
        "Проект",
        "Этап",
        "Веха / часть",
        "План",
        "Индивидуальный срок",
        "Факт / событие",
        "Отклонение, дней",
        "Вес",
        "Сумма, ₽",
        "Статус",
        "Дата выплаты",
    )
    rows, totals = [], []
    for value in results.values():
        case, employee = value["case"], value["result"]["employees"][0]
        if case["project_code"] != project_code or case["stage"] != stage:
            continue
        for item in [*employee["milestones"], employee["final"], *employee["payouts"]]:
            rows.append(
                [
                    employee["tab_no"],
                    employee["fio"],
                    case["project"],
                    stage,
                    item.get("name", item.get("part", employee["final"]["kind"])),
                    item.get("plan"),
                    item.get("individual_plan"),
                    item.get("fact", item.get("event_date")),
                    item.get("deviation_days"),
                    item.get("weight"),
                    money(item["amount"]) if item.get("amount") is not None else None,
                    item["calc_status"],
                    item.get("pay_date"),
                ]
            )
        totals.append(
            [
                employee["tab_no"],
                employee["fio"],
                money(employee["pb"]) if employee["pb"] is not None else None,
                employee["k_fin"],
                employee["k_bs"],
                money(employee["interest"]["total"]),
                money(employee["total_payable"]),
            ]
        )
    return excel_bytes(
        {
            "калькулятор": (columns, rows),
            "Итоги": (
                (
                    "Табельный номер",
                    "ФИО",
                    "ПБ, ₽",
                    "К фин",
                    "Множитель БС",
                    "Проценты, ₽",
                    "Итого, ₽",
                ),
                totals,
            ),
            "Сведения": (
                ("Статус выписки",),
                [
                    [
                        "Расчёт предварительный"
                        if preliminary
                        else "Выписка на дату закрытия квартала; будущие части предварительные"
                    ]
                ],
            ),
        }
    )


def problems_file(problems):
    columns = ("Код", "Табельный номер", "Проект", "Веха", "Блокирует", "Описание")
    return csv_bytes(
        columns,
        [
            [
                p["code"],
                p.get("tab_no"),
                p.get("project_code"),
                p.get("milestone"),
                "Да" if p["blocking"] else "Нет",
                p["message"],
            ]
            for p in problems
        ],
    )


def summary_file(entries, previous=()):
    def aggregate(rows):
        totals = {}
        for row in rows:
            if row["calc_status"] == "рассчитано":
                key = (row["tab_no"], row["fio"], row["project"], row["stage"])
                totals[key] = totals.get(key, 0) + row["amount"]
        return totals

    current, old = aggregate(entries), aggregate(previous)
    rows = [
        [
            *key,
            money(current.get(key, 0)),
            money(old.get(key, 0)),
            money(current.get(key, 0) - old.get(key, 0)),
        ]
        for key in sorted(current.keys() | old.keys())
    ]
    return excel_bytes(
        {
            "Сводная": (
                (
                    "Табельный номер",
                    "ФИО",
                    "Проект",
                    "Этап",
                    "Текущий квартал, ₽",
                    "Предыдущий квартал, ₽",
                    "Разница, ₽",
                ),
                rows,
            )
        }
    )
